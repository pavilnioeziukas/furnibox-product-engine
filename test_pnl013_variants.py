from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

from approved_bom_replay import stage_approved_v10_input
from calculator_settings import validate
from price_calculators import source
from reform_so_line_prices import apply_missing_pnl013_rule, key
from so_pricing_rules import PricingRule
from test_approved_bom_replay import base_rows
from unified_calculator_pricing import recipes


def test_staged_workbook_has_all_panel_variants_and_parts(tmp_path):
    original = tmp_path / 'source.xlsx'
    staged = tmp_path / 'staged.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'BOM - Input'
    sheet.cell(8, 3, 'BOM SKU Code')
    for number, row in enumerate(base_rows(), 9):
        for column, value in enumerate(row, 1):
            if value is not None:
                sheet.cell(number, column, value)
    workbook.save(original)
    workbook.close()
    stage_approved_v10_input(original, staged)
    workbook = load_workbook(staged, data_only=True)
    rows = {row[2]: row for row in workbook.active.iter_rows(min_row=9, max_col=133, values_only=True)}
    catalog = {row['sku']: row for row in source('panel')['rows']}
    registry = recipes(validate({}))
    numbers = []
    for cabinet, color in [('CAB01', 'WW'), ('CAB02', 'BB'), ('CAB03', 'NO')]:
        sku = f'EUB-C-{cabinet}-PNL013'
        detail = f'EU-PNL-2000x150-{color}'
        assert rows[sku][13:17] == (detail, 1, 'PANEL', 0)
        assert catalog[sku]['detail'] == detail
        assert (catalog[sku]['length'], catalog[sku]['width']) == (2000, 150)
        assert registry[detail.casefold()]['total'] == pytest.approx(catalog[sku]['sourceK'])
        assert registry[sku.casefold()]['total'] == pytest.approx(catalog[sku]['sourceK'] + .9 + 1.5)
        numbers.append(rows[sku][1])
    assert len(set(numbers)) == 3
    workbook.close()
    repeated = tmp_path / 'repeated.xlsx'
    changes = stage_approved_v10_input(staged, repeated)
    assert all(row['action'] == 'ALREADY_MATCHING' for row in changes)


def test_each_panel_inherits_its_color_tariff_and_keeps_explicit_override():
    rules = {}
    for index, cabinet in enumerate(['CAB01', 'CAB02', 'CAB03'], 1):
        sku = f'EUB-C-{cabinet}-PNL011'
        rules[key(sku)] = PricingRule(sku, '', '', '', index, 0, 0, 0, 0, 0)
    result = apply_missing_pnl013_rule(rules)
    for cabinet in ['CAB01', 'CAB02', 'CAB03']:
        target = f'EUB-C-{cabinet}-PNL013'
        assert result[key(target)] == replace(rules[key(f'EUB-C-{cabinet}-PNL011')], sku=target)
    target = 'EUB-C-CAB01-PNL013'
    override = PricingRule(target, '', '', '', 99, 0, 0, 0, 0, 0)
    assert apply_missing_pnl013_rule({**rules, key(target): override})[key(target)] == override
