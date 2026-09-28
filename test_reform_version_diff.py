import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from flask import Flask
from openpyxl import Workbook, load_workbook
from reform_version_diff import compare_versions, read_version
from webapp.reform_versions import versions

HEADERS = ['BOM SKU Code', 'Name', 'Part 1 Code', 'Part 1 Qty', 'Part 2 Code', 'Part 2 Qty']


class VersionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write(self, name, rows, headers=HEADERS, sheet='BOM - Input'):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(headers)
        for row in rows:
            ws.append(row)
        path = self.root / name
        wb.save(path)
        return path

    def test_reorder_and_repeated_component_have_no_false_changes(self):
        a = self.write('a.xlsx', [['A','Name','X',1,'X',2],['B','Bee','Y',1]])
        b = self.write('b.xlsx', [['B','Bee','Y',1],['A','Name','X',3]])
        result = compare_versions(a, b)
        self.assertEqual(result['changes'], [])
        self.assertEqual(result['counts']['Nepakitęs BOM'], 2)

    def test_all_change_kinds(self):
        a = self.write('a.xlsx', [['A','Old','X',1,'Y',2],['B','Gone','Q',1]])
        b = self.write('b.xlsx', [['A','New','X',3,'Z',4],['C','Added','R',1]])
        result = compare_versions(a, b)
        self.assertEqual(result['counts'], {'Pakeistas BOM':1,'Pašalintas BOM':1,'Naujas BOM':1})
        self.assertEqual(result['added_skus'], ['C','R','Z'])
        self.assertEqual(result['removed_skus'], ['B','Q','Y'])
        quantity = next(r for r in result['changes'] if r['status'] == 'Pakeistas kiekis')
        self.assertEqual((quantity['before'], quantity['after']), ('1','3'))
        self.assertTrue(any(r['status'] == 'Pakeistas laukas' for r in result['changes']))

    def test_bad_quantities_duplicate_and_empty_are_rejected(self):
        for value in (None, 0, -1, 'NaN', '=1+1'):
            with self.subTest(value=value):
                p = self.write('bad.xlsx', [['A','Name','X',value]])
                with self.assertRaises(ValueError):
                    read_version(p)
        for rows in ([], [['A','Name','X',1],['A','Name','X',1]]):
            with self.assertRaises(ValueError):
                read_version(self.write('bad.xlsx', rows))

    def test_decimal_precision_and_dynamic_slots(self):
        headers = ['BOM SKU Code','Part 37 Code','Part 37 Qty','Part 38 Code','Part 38 Qty']
        a = self.write('a.xlsx', [['A','X',0.1,'X',0.2]], headers)
        b = self.write('b.xlsx', [['A','X',0.3]], headers)
        self.assertEqual(compare_versions(a,b)['changes'], [])

    def test_vertical_supported(self):
        headers = ['BOM SKU Code','Part Code','Qty','Part Group']
        p = self.write('vertical.xlsx', [['A','X',2,'Hardware'],['A','Y',3,'Hardware']], headers, 'BOM VERTICAL')
        self.assertEqual(len(read_version(p)[1]['A']['parts']), 2)

    def test_real_full_db_headers_and_row_numbers(self):
        headers = ['Row ID','SKU Code','SKU Name','Part #','Part Column','Part Code','Part Qty.','Part Group']
        p = self.write('db.xlsx', [[1,'A','Alpha',1,14,'X',2,'Hardware'],[2,'A','Alpha',2,18,'Y',3,'Hardware']], headers, 'BOM - Full DB')
        sheet, products = read_version(p)
        self.assertEqual(sheet, 'BOM - Full DB')
        self.assertEqual(products['A']['fields'], {'SKU Name':'Alpha'})
        self.assertEqual(products['A']['parts']['X']['fields'], {'Group':'Hardware'})
        self.assertEqual(len(products['A']['parts']), 2)

    def test_input_preferred_to_stale_derived_database(self):
        p = self.write('source.xlsx', [['A','Alpha','X',2]])
        wb = load_workbook(p)
        stale = wb.create_sheet('BOM - Full DB', 0)
        stale.append(['SKU Code','Part Code','Part Qty.'])
        stale.append(['OLD','REMOVED',99])
        wb.save(p)
        sheet, products = read_version(p)
        self.assertEqual(sheet, 'BOM - Input')
        self.assertEqual(list(products), ['A'])

    def test_routes_export_and_input_validation(self):
        a = self.write('old.xlsx', [['A','Old','X',1]])
        b = self.write('new.xlsx', [['A','=unsafe','X',2]])
        # Literal imported text must survive export as text, never as a formula.
        wb = load_workbook(b)
        wb.active['B2'].data_type = 's'
        wb.save(b)
        app = Flask(__name__, template_folder=str(Path(__file__).parent/'webapp/templates'))
        app.secret_key = 'test'
        app.config.update(TESTING=True, REFORM_VERSION_UPLOADS=self.root, REFORM_VERSION_REPORTS=self.root/'reports')
        app.register_blueprint(versions)
        app.add_url_rule('/', 'index', lambda: '')
        app.add_url_rule('/reports', 'reports.index', lambda: '')
        app.add_url_rule('/logout', 'logout', lambda: '')
        @app.context_processor
        def context():
            return dict(product_engine=SimpleNamespace(show_pricing_nav=False, app_name='Test', brand_mark='F', brand_name='Furnibox'), bootstrap_manager_enabled=False)
        client = app.test_client()
        self.assertEqual(client.get('/reform-versions').status_code, 200)
        with client.session_transaction() as session:
            token = session['version_diff_csrf']
        self.assertEqual(client.post('/reform-versions', data={}).status_code, 400)
        invalid = client.post('/reform-versions', data=dict(csrf_token=token, old='../old.xlsx', new=b.name))
        self.assertIn('Pasirinkite du'.encode(), invalid.data)
        response = client.post('/reform-versions', data=dict(csrf_token=token, old=a.name, new=b.name))
        self.assertEqual(response.status_code, 302)
        location = response.headers['Location']
        self.assertIn('Kas pasikeitė'.encode(), client.get(location).data)
        self.assertIn('Ankstesnės ataskaitos'.encode(), client.get('/reform-versions').data)
        exported = client.get(location+'?format=xlsx')
        self.assertEqual(exported.status_code, 200)
        wb = load_workbook(io.BytesIO(exported.data), data_only=False)
        cells = [c for row in wb['Pakeitimai'] for c in row if c.value == '=unsafe']
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, 's')
        self.assertEqual(client.get('/reform-versions/not-a-report').status_code, 404)


if __name__ == '__main__':
    unittest.main()
