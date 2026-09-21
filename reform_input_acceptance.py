"""Read-only checks of approved BOM corrections in a supplied Reform workbook.

This deliberately does not rewrite a supplier workbook or alter Odoo. Missing
and conflicting rows are distinguished so a future import can decide whether
an approved correction may be applied or needs human review.
"""

from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from confirmed_purchased_products import (
    APPROVED_SUPPLIER_ALIASES,
    CONFIRMED_PURCHASED_NON_BOM_SKUS,
)
from reform_map import detect_columns, find_header, find_sheet


BASE = Path(__file__).resolve().parent
EXPECTATIONS = BASE / "manifest" / "approved_reform_bom_expectations.json"


def rule_snapshot() -> dict:
    """Show which persistent correction registries this run will use."""
    manifest = BASE / "manifest"
    review = json.loads((manifest / "pricing_review_corrections.json").read_text(encoding="utf-8"))
    allocation = json.loads((manifest / "apack_hrd_allocation_rules.json").read_text(encoding="utf-8"))
    prices = json.loads((manifest / "purchase_price_adjustments.json").read_text(encoding="utf-8"))
    return {
        "purchased_non_bom_skus": len(CONFIRMED_PURCHASED_NON_BOM_SKUS),
        "approved_supplier_aliases": len(APPROVED_SUPPLIER_ALIASES),
        "reviewed_pricing_corrections": len(review["corrections"]),
        "apack_hrd_family_rules": len(allocation),
        "approved_purchase_prices": len(prices["adjustments"]),
    }


def canonical(value: object) -> str:
    return str(value or "").strip().upper()


def read_selected_boms(
    path: Path, expected_skus: set[str], purchased_skus: set[str] | None = None,
) -> tuple[dict, list[dict], list[dict]]:
    """Read only approved SKUs, using named Part Code / Qty columns."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = find_sheet(workbook)
        header_row, headers = find_header(sheet)
        sku_column, parts = detect_columns(headers)
        result = {}
        issues = []
        purchased_rows = []
        purchased_skus = purchased_skus or set()
        seen_boms = set()
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            sku = canonical(row[sku_column] if sku_column < len(row) else None)
            if not sku:
                continue
            if sku in seen_boms:
                issues.append({"type": "DUPLICATE_SOURCE_BOM", "sku": sku, "row": row_number})
            seen_boms.add(sku)
            if sku in purchased_skus:
                purchased_rows.append({"sku": sku, "source_row": row_number, "action": "IGNORE_SOURCE_BOM"})
            if sku in result:
                continue
            components = defaultdict(float)
            seen_parts = set()
            for part in parts:
                component = canonical(row[part.code_index] if part.code_index < len(row) else None)
                if not component:
                    continue
                if component in seen_parts:
                    issues.append({
                        "type": "DUPLICATE_SOURCE_COMPONENT", "sku": sku,
                        "component": component, "row": row_number,
                    })
                seen_parts.add(component)
                quantity = row[part.qty_index] if part.qty_index < len(row) else None
                if isinstance(quantity, bool) or not isinstance(quantity, (float, int)) or not math.isfinite(quantity) or quantity <= 0:
                    issues.append({
                        "type": "INVALID_QUANTITY", "sku": sku,
                        "component": component, "row": row_number,
                    })
                    continue
                components[component] += float(quantity)
            if sku in expected_skus:
                result[sku] = dict(components)
        return result, issues, purchased_rows
    finally:
        workbook.close()


def compare_boms(actual: dict, expected: dict) -> list[dict]:
    """Separate absence from disagreement; never silently accept a mismatch."""
    issues = []
    for original_sku, components in expected.items():
        sku = canonical(original_sku)
        wanted = {canonical(part): float(qty) for part, qty in components.items()}
        found = actual.get(sku)
        if found is None:
            issues.append({"type": "MISSING_APPROVED_BOM", "sku": sku, "expected": wanted})
            continue
        found = {canonical(part): float(qty) for part, qty in found.items()}
        for part in sorted(set(wanted) | set(found)):
            expected_qty, actual_qty = wanted.get(part, 0.0), found.get(part, 0.0)
            if not math.isclose(expected_qty, actual_qty, abs_tol=1e-9):
                issues.append({
                    "type": "CONFLICTING_BOM_COMPONENT", "sku": sku,
                    "component": part, "expected_qty": expected_qty,
                    "actual_qty": actual_qty,
                })
    return issues


def audit_input(path: Path, expectations_path: Path = EXPECTATIONS) -> dict:
    rules = json.loads(expectations_path.read_text(encoding="utf-8"))
    expected = rules["boms"]
    try:
        actual, read_issues, purchased_rows = read_selected_boms(
            path,
            {canonical(sku) for sku in expected},
            {canonical(sku) for sku in CONFIRMED_PURCHASED_NON_BOM_SKUS},
        )
    except (KeyError, ValueError) as error:
        actual, purchased_rows = {}, []
        read_issues = [{"type": "INVALID_REFORM_INPUT_FORMAT", "detail": str(error)}]
    format_error = any(issue["type"] == "INVALID_REFORM_INPUT_FORMAT" for issue in read_issues)
    issues = read_issues + ([] if format_error else compare_boms(actual, expected))
    return {
        "rules_version": rules["version"],
        "source_file": path.name,
        "checked_boms": len(expected),
        "matched_boms": sum(
            1 for sku in expected
            if canonical(sku) in actual
            and not any(issue.get("sku") == canonical(sku) for issue in issues)
        ),
        "persistent_rule_coverage": rule_snapshot(),
        "validated_at_input": [
            "approved_bom_components",
            "source_bom_duplicates_and_quantities",
            "purchased_skus_with_source_bom_listed_for_exclusion",
        ],
        "requires_downstream_price_validation": [
            "supplier_aliases",
            "pricing_review_corrections",
            "apack_hrd_allocation",
            "purchase_price_adjustments",
            "calculator_and_final_price_rules",
        ],
        "purchased_source_boms_to_ignore": purchased_rows,
        "issues": issues,
        "status": "PASS" if not issues else "REVIEW",
    }


def audit_generated_boms(path: Path, expectations_path: Path = EXPECTATIONS) -> dict:
    """Check approved -A content in the calculated component-cost trace."""
    rules = json.loads(expectations_path.read_text(encoding="utf-8"))
    expected = rules.get("post_generation_boms", {})
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "BOM COMPONENT COSTS" not in workbook.sheetnames:
            return {
                "status": "REVIEW",
                "issues": [{"type": "MISSING_COMPONENT_TRACE_SHEET"}],
            }
        sheet = workbook["BOM COMPONENT COSTS"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return {
                "status": "REVIEW",
                "issues": [{"type": "EMPTY_COMPONENT_TRACE_SHEET"}],
            }
        columns = {str(value or "").strip(): index for index, value in enumerate(header)}
        required = {
            "Top BOM SKU", "Level II SKU", "Purchased Component SKU",
            "Component Qty in Level II",
        }
        if not required.issubset(columns):
            return {
                "status": "REVIEW",
                "issues": [{"type": "INVALID_COMPONENT_TRACE_FORMAT", "missing_columns": sorted(required - set(columns))}],
            }
        parents = defaultdict(lambda: defaultdict(dict))
        issues = []
        expected_keys = {canonical(sku) for sku in expected}
        for row in rows:
            level_two = canonical(row[columns["Level II SKU"]])
            if level_two not in expected_keys:
                continue
            top = canonical(row[columns["Top BOM SKU"]])
            component = canonical(row[columns["Purchased Component SKU"]])
            qty = row[columns["Component Qty in Level II"]]
            if component in parents[level_two][top]:
                issues.append({"type": "DUPLICATE_GENERATED_COMPONENT", "sku": level_two, "top_bom_sku": top, "component": component})
            elif isinstance(qty, (int, float)) and not isinstance(qty, bool) and math.isfinite(qty) and qty > 0:
                parents[level_two][top][component] = float(qty)
            else:
                issues.append({"type": "INVALID_GENERATED_QUANTITY", "sku": level_two, "top_bom_sku": top, "component": component})
        checked_parents = 0
        for sku, wanted in expected.items():
            normalized = canonical(sku)
            if not parents[normalized]:
                issues.append({"type": "MISSING_GENERATED_BOM_TRACE", "sku": normalized})
                continue
            for top, actual in parents[normalized].items():
                checked_parents += 1
                for issue in compare_boms({normalized: actual}, {normalized: wanted}):
                    issues.append({**issue, "top_bom_sku": top})
        return {
            "status": "PASS" if not issues else "REVIEW",
            "checked_parent_boms": checked_parents,
            "issues": issues,
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = audit_input(args.input)
    content = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(content + "\n", encoding="utf-8")
    print(content)
    return 0 if report["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
