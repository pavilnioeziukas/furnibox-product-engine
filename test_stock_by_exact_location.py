import ast
from pathlib import Path

import openpyxl
import pytest

from run_stock_by_exact_location import StockByExactLocationReport


class Client:
    def __init__(self):
        self.calls = []

    def search(self, model, domain, **kwargs):
        self.calls.append((model, domain))
        return [1, 2, 3, 4, 5]

    def read(self, model, ids, fields=None):
        if model == "stock.location":
            return [{"id": i, "complete_name": {8: "WH/Stock", 9: "WH/Stock/A1"}[i]} for i in ids]
        return [
            {"id": 1, "product_id": [10, "A"], "location_id": [8, "Stock"], "quantity": 5, "reserved_quantity": 2},
            {"id": 2, "product_id": [10, "A"], "location_id": [8, "Stock"], "quantity": 3, "reserved_quantity": 1},
            {"id": 3, "product_id": [10, "A"], "location_id": [9, "A1"], "quantity": -2, "reserved_quantity": 0},
            {"id": 4, "product_id": [11, "B"], "location_id": [9, "A1"], "quantity": 0, "reserved_quantity": 4},
            {"id": 5, "product_id": [99, "Excluded"], "location_id": [9, "A1"], "quantity": 100, "reserved_quantity": 0},
        ]


def products():
    # Duplicate SKU codes still represent different Odoo products.
    return [{"id": 11, "default_code": "SKU", "display_name": "B"},
            {"id": 10, "default_code": "SKU", "display_name": "A"}]


def test_exact_grouping_keeps_child_separate_and_sums_quant_dimensions():
    client = Client()
    report = StockByExactLocationReport(client)
    balances, locations = report.load_exact_balances({10, 11})
    assert balances == {
        (10, 8): {"on_hand": 8, "reserved": 3},
        (10, 9): {"on_hand": -2, "reserved": 0},
        (11, 9): {"on_hand": 0, "reserved": 4},
    }
    assert client.calls[0][1][0] == ("location_id.usage", "=", "internal")
    assert "child_of" not in str(client.calls)
    rows = report.build_exact_rows(products(), set(), {}, balances, locations)
    assert [(r["Product"], r["Location"], r["Available"]) for r in rows] == [
        ("A", "WH/Stock", 5), ("A", "WH/Stock/A1", -2), ("B", "WH/Stock/A1", -4)]
    assert not any(key.startswith(("WH ", "C ")) for row in rows for key in row)


def test_zero_totals_removed_and_missing_path_fails():
    report = StockByExactLocationReport(Client())
    assert report.build_exact_rows(products(), set(), {}, {(10, 8): {"on_hand": 0, "reserved": 0}}, {}) == []
    with pytest.raises(ValueError, match="lokacijos 8"):
        report.build_exact_rows(products(), set(), {}, {(10, 8): {"on_hand": 1, "reserved": 0}}, {})


def test_run_and_workbook_do_not_repeat_global_po_per_location(monkeypatch, tmp_path):
    report = StockByExactLocationReport(Client())
    monkeypatch.setattr(report, "load_products", products)
    monkeypatch.setattr(report, "load_categories", lambda p: {})
    monkeypatch.setattr(report, "find_buy_product_ids", lambda **kwargs: set())
    monkeypatch.setattr(report, "load_last_purchases", lambda ids: {})
    monkeypatch.setattr(report, "load_active_po_remaining", lambda p: {"SKU": 7})
    path = report.run(str(tmp_path))
    workbook = openpyxl.load_workbook(path, data_only=True)
    stock = workbook["Stock by Location"]
    summary = workbook["SKU Summary"]
    assert stock.max_row == 4
    assert report.INCOMING_PO_COLUMN not in [c.value for c in stock[1]]
    assert list(summary.values)[1] == ("SKU", 6, 7, -1, 7)
    assert stock.freeze_panes == "A2"
    assert stock.auto_filter.ref == stock.dimensions
    workbook.close()


def test_empty_workbook_keeps_headers(tmp_path):
    path = StockByExactLocationReport.export_exact_excel([], [], str(tmp_path))
    workbook = openpyxl.load_workbook(path)
    assert workbook["Stock by Location"].max_row == 1
    assert workbook["SKU Summary"]["A1"].value == "SKU"
    workbook.close()


def test_web_action_is_separate_and_uses_job_output_directory():
    tree = ast.parse(Path("webapp/app.py").read_text(encoding="utf-8"))
    action = next(node for node in ast.walk(tree) if isinstance(node, ast.Dict)
                  and any(isinstance(k, ast.Constant) and k.value == "stock_by_exact_location" for k in node.keys))
    index = next(i for i, key in enumerate(action.keys) if isinstance(key, ast.Constant) and key.value == "stock_by_exact_location")
    settings = ast.literal_eval(action.values[index])
    assert settings["module"] == "run_stock_by_exact_location"
    assert settings["args"] == ["--output-dir", "{output_dir}"]
    assert settings["collect_changed_outputs"] is False
    assert '"module": "run_stock_by_location"' in Path("webapp/app.py").read_text(encoding="utf-8")
