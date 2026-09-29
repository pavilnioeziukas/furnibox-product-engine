import io,copy,json
from openpyxl import load_workbook
from test_reform_workspace import setup,post,sample
from test_reform_excel import upload_book
from webapp.reform_workspace import db,baseline,pack,unpack


def export(client):
    # Recreate a saved pre-vertical export to verify old files remain importable.
    from webapp.reform_excel import make_book
    r=client.get('/reform/excel/download?scope=all&layout=legacy')
    assert r.status_code==200
    token=load_workbook(io.BytesIO(r.data))['_Exchange']['B1'].value
    with client.application.app_context(),db() as conn:
        record=unpack(conn.execute('SELECT payload FROM file_exports WHERE id=?',(token,)).fetchone()[0])
        record.update(map_format=2,bom_format=0)
        conn.execute('UPDATE file_exports SET payload=? WHERE id=?',(pack(record),token))
    return load_workbook(make_book(token,record))


def test_map_edit_add_non_bom_and_supplier_reference(setup):
    app,client=setup
    with app.app_context(),db() as conn:
        conn.execute('INSERT INTO state VALUES (3,?)',(pack({'PANEL-01':['SUP-PANEL']}),))
    b=export(client)
    assert b.sheetnames[:2]==['BOM Map','Non-BOM']
    assert b['Components'].sheet_state=='veryHidden'
    assert b['Products'].max_row==2 and b['Non-BOM'].max_row==3
    h={c.value:c.column for c in b['BOM Map'][1]}; s=b['BOM Map']
    assert s.cell(2,h['Supplier Code']).value=='SUP-PANEL'
    assert 'No changes detected' in upload_book(client,b).text
    s.cell(2,h['Supplier Code'],'OTHER')
    assert 'reference field' in upload_book(client,b).text
    s.cell(2,h['Supplier Code'],'SUP-PANEL')
    s.cell(2,h['Level II Qty'],3)
    assert 'Excel changes saved' in upload_book(client,b).text
    with app.app_context(),db() as conn:
        work=unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert work['target']['boms']['1']['components'][0]['quantity']==3
    b=export(client); s=b['BOM Map']; row=[None]*s.max_column
    row[h['Parent BOM SKU']-1]='CAB-01'; row[h['Purchased Component SKU']-1]='PANEL-01'
    row[h['Component Qty in Level II']-1]=5; row[h['Action']-1]='KEEP'; s.append(row)
    assert 'Excel changes saved' in upload_book(client,b).text


def test_nested_map_normalization_and_repeated_line_conflict(setup):
    app,client=setup
    with app.app_context(),db() as conn:
        data=baseline(conn)
        data['products']['SUB']={**data['products']['CAB-01'],'sku':'SUB'}
        data['boms']['1']['components'][0].update(sku='SUB',quantity=4)
        data['boms']['1']['quantity']=2
        data['boms']['2']={'id':'2','sku':'SUB','code':'Sub','quantity':2,'active':True,'uom':'vnt.','uom_id':1,'components':[{'id':'20','sku':'PANEL-01','quantity':6,'uom':'vnt.','uom_id':1}]}
        conn.execute('UPDATE state SET payload=? WHERE id=1',(pack(data),))
    b=export(client);s=b['BOM Map'];h={c.value:c.column for c in s[1]}
    assert s.cell(2,h['Total Qty in Top BOM']).value==6
    assert 'No changes detected' in upload_book(client,b).text
    s.cell(2,h['Component Qty in Level II'],8)
    last=s.max_row
    s.cell(last,h['Level II Qty'],9)
    assert 'conflicting edits' in upload_book(client,b).text
    s.cell(last,h['Level II Qty'],6)
    assert 'Excel changes saved' in upload_book(client,b).text
    with app.app_context(),db() as conn:
        work=unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert work['target']['boms']['2']['components'][0]['quantity']==8


def test_map_missing_rows_and_non_bom_changes(setup):
    app,client=setup
    b=export(client);b['BOM Map'].delete_rows(2)
    assert 'rows are missing' in upload_book(client,b).text
    b=export(client);b['Non-BOM']['B2']='Updated component'
    assert 'Excel changes saved' in upload_book(client,b).text


def test_supplier_reference_is_admin_only_and_preserves_drafts(setup):
    app,client=setup
    with client.session_transaction() as sess: csrf=sess['reform_csrf']
    assert client.post('/reform/supplier-reference',data={'csrf':csrf,'file':(io.BytesIO(b'{}'),'codes.json')}).status_code==403
    app.config['REFORM_ADMIN_ENABLED']=True
    with client.session_transaction() as sess:
        sess.pop('reform_user');sess['authenticated']=True
    r=client.post('/reform/supplier-reference',data={'csrf':csrf,'file':(io.BytesIO(b'{"PANEL-01":["SUP-1"]}'),'codes.json')})
    assert r.status_code==302
    with app.app_context(),db() as conn:
        assert baseline(conn)==sample()
