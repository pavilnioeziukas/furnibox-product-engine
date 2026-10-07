"""Exercise complete Excel round trips through the authenticated application."""
import io
import copy
import json
import re

from openpyxl import load_workbook
from test_reform_workspace import setup, post, sample
from webapp.reform_workspace import db, baseline, unpack


def export_book(client, product='CAB-01'):
    page = client.get('/reform/versions')
    if 'Open a version draft' in page.text:
        post(client, 'versions/save', number='v1.0', description='Excel test version', revision='0')
    response = client.get('/reform/excel/download', query_string={'product': product, 'layout': 'legacy'})
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.data))


def test_excel_exchange_requires_open_version(setup):
    app, client = setup
    response = client.get('/reform/excel/download', follow_redirects=True)
    assert 'Open a version draft before downloading or uploading changes.' in response.text
    assert 'Open a version first' in client.get('/reform/?view=files').text
    post(client, 'versions/save', number='v1.0', description='Catalogue update', revision='0')
    page = client.get('/reform/?view=files')
    assert 'Open version:</strong> v1.0' in page.text
    assert client.get('/reform/excel/download').status_code == 200


def upload_book(client, wb):
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    with client.session_transaction() as sess:
        csrf = sess['reform_csrf']
    return client.post('/reform/excel/upload', data={'csrf': csrf, 'file': (data, 'Reform.xlsx')}, follow_redirects=True)


def test_unchanged_and_changed_roundtrip(setup):
    app, client = setup
    book = export_book(client)
    assert 'No changes detected' in upload_book(client, book).text
    book['Components']['E2'] = 7
    response = upload_book(client, book)
    assert 'Excel changes saved' in response.text
    assert 'PANEL-01 — 7.0' in re.sub(r'<[^>]+>', '', response.text)
    assert 'Current' in response.text and 'Proposed' in response.text
    assert 'Your draft has changed' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        assert baseline(conn) == sample()
        assert conn.execute('SELECT count(*) FROM submissions').fetchone()[0] == 0
    assert 'submitted to Furnibox' in post(client, 'submit', revision='1', confirm='yes').text


def test_missing_rows_and_explicit_removal(setup):
    app, client = setup
    book = export_book(client)
    book['Components'].delete_rows(3)
    assert 'exported rows are missing' in upload_book(client, book).text
    book = export_book(client)
    book['Components']['C3'] = 'REMOVE'
    assert 'Excel changes saved' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        work = unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert len(work['target']['boms']['1']['components']) == 1


def test_zero_quantity_removes_existing_component(setup):
    app, client = setup
    book = export_book(client)
    book['Components']['E3'] = 0
    assert 'Excel changes saved' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        work = unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert [line['id'] for line in work['target']['boms']['1']['components']] == ['11']


def test_new_cards_new_bom_and_new_component(setup):
    app, client = setup
    with app.app_context(), db() as conn:
        data = sample()
        for p in data['products'].values():
            p.update(category='All / Parts', category_id=7)
        conn.execute('UPDATE state SET payload=? WHERE id=1', (json.dumps(data),))
    book = export_book(client, '')
    book['Products'].append(['NEW-CAB', 'New cabinet', 'All / Parts', 'vnt.'])
    book['Products'].append(['NEW-PART', 'New panel', 'All / Parts', 'vnt.'])
    book['BOMs'].append(['NEW-1', 'NEW-CAB', 'New cabinet BOM', 1])
    book['Components'].append(['NEW-1', None, 'KEEP', 'NEW-PART', 3])
    assert 'Excel changes saved' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        work = unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert work['target']['products']['NEW-PART']['category_id'] == 7
        assert len(work['target']['boms']) == 2
        assert baseline(conn)['products'].get('NEW-CAB') is None


def test_stale_source_and_other_user(setup):
    app, client = setup
    book = export_book(client)
    with client.session_transaction() as sess:
        sess['reform_user'] = 'other'
    assert 'only its owner can edit it' in upload_book(client, book).text
    with client.session_transaction() as sess:
        sess['reform_user'] = 'paul'
    with app.app_context(), db() as conn:
        data = sample()
        data['captured_at'] = 'later'
        conn.execute('UPDATE state SET payload=? WHERE id=1', (json.dumps(data),))
    assert 'Source data has changed' in upload_book(client, book).text


def test_invalid_workbooks_are_atomic(setup):
    app, client = setup
    for value, error in [(-1, 'greater than zero'), ('=1+1', 'replace formulas')]:
        book = export_book(client)
        book['Components']['E2'] = value
        assert error in upload_book(client, book).text
    book = export_book(client)
    book['Components']['D2'] = 'APACK-PRIVATE'
    assert 'component does not exist' in upload_book(client, book).text
    book = export_book(client)
    book['Components']['B2'] = 'forged'
    assert 'Line ID is invalid' in upload_book(client, book).text
    book = export_book(client)
    book['Components']['D2'] = 'CAB-01'
    assert 'cycle' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        assert conn.execute('SELECT count(*) FROM drafts').fetchone()[0] == 0


def test_nested_export_and_row_reordering(setup):
    app, client = setup
    with app.app_context(), db() as conn:
        data = sample()
        data['boms']['2'] = {**copy.deepcopy(data['boms']['1']), 'id': '2', 'sku': 'PANEL-01',
                            'components': [{'id': '21', 'sku': 'HINGE-01', 'quantity': 1, 'uom': 'vnt.', 'uom_id': 1}]}
        conn.execute('UPDATE state SET payload=? WHERE id=1', (json.dumps(data),))
    book = export_book(client)
    assert book['BOMs'].max_row == 3
    sheet = book['Components']
    values = list(sheet.values)[1:]
    sheet.delete_rows(2, len(values))
    for row in reversed(values):
        sheet.append(row)
    assert 'No changes detected' in upload_book(client, book).text


def test_upload_requires_csrf_and_file_page_renders(setup):
    app, client = setup
    assert client.post('/reform/excel/upload', data={}).status_code == 400
    assert 'Open a version first' in client.get('/reform/?view=files').text
    post(client, 'versions/save', number='v1.0', description='Excel exchange', revision='0')
    assert 'Upload and compare' in client.get('/reform/?view=files').text

def test_full_catalogue_roundtrip_and_read_only_bom(setup):
    app, client = setup
    with app.app_context(), db() as conn:
        data = baseline(conn)
        data['products']['UNUSED'] = dict(data['products']['PANEL-01'], sku='UNUSED', name='Unused card')
        data['boms']['2'] = copy.deepcopy(data['boms']['1'])
        data['boms']['2'].update(id='2', sku='UNUSED', read_only='Internal assembly')
        conn.execute('UPDATE state SET payload=? WHERE id=1', (json.dumps(data),))
    post(client, 'versions/save', number='v1.0', description='Full catalogue', revision='0')
    response = client.get('/reform/excel/download?scope=all&layout=legacy')
    book = load_workbook(io.BytesIO(response.data))
    assert book['Products'].max_row == 5
    assert book['BOMs'].max_row == 3
    assert book['Read-only']['B2'].value == '2'
    assert 'No changes detected' in upload_book(client, book).text
    book['Components']['E4'] = 9
    assert 'read-only' in upload_book(client, book).text
    book['Components']['E4'] = 2
    book['Components']['E2'] = 7
    assert 'Excel changes saved' in upload_book(client, book).text
    with app.app_context(), db() as conn:
        work = unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])
        assert work['target']['boms']['2'] == data['boms']['2']

