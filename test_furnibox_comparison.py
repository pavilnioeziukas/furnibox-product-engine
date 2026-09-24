import csv
import io
import pytest
from flask import Flask
import furnibox_comparison as model
from webapp.furnibox_comparison import comparison


def test_catalog_and_tamara_coefficients():
    rows=model.source()['rows']
    assert len(rows)==len({r['sku'] for r in rows})==1884
    assert sum(r['coefficient']==4 for r in rows)==48
    assert sum(r['coefficient']==2.2 for r in rows)==30
    assert not any('U732ST9' in r['sku'] for r in rows)


@pytest.mark.parametrize('sku,edge',[('EU-CUTBACK-600x400-WW',0),('EU-HALFBACK-600x400-WW',0),('EU-BOT-600x400-WW',2),('EU-STIFF-600x400-WW',2)])
def test_material_and_perimeter(sku,edge):
    row=dict(model.source()['rows'][2],sku=sku,length=600,width=400,color='WW',coefficient=2)
    v=model.calculate(row,model.source()['rates'])
    expected=.24*(2.85 if 'BACK' in sku else 8.72035)*1.1+edge*1.1*.2
    assert v['material']==pytest.approx(expected)
    assert v['calculated_purchase']==pytest.approx(expected*2)
    assert v['furnibox_coefficient']==pytest.approx(row['reform']/(expected*2))


def test_persistence_and_override_isolation(tmp_path):
    rows=model.source()['rows'];row=rows[2]
    rates,before=model.results(tmp_path)
    model.save_coefficient(tmp_path,row['sku'],'2,5')
    _,after=model.results(tmp_path)
    assert after[2]['coefficient']==2.5
    assert after[2]['calculated_purchase']==pytest.approx(before[2]['material']*2.5)
    assert before[3:]==after[3:]
    assert after[2]['purchase']==before[2]['purchase']
    assert after[2]['difference']==before[2]['difference']
    model.save_rates(tmp_path,dict(rates,BACK=3))
    _,updated=model.results(tmp_path)
    assert updated[2]['coefficient']==2.5
    assert updated[2]['material']>after[2]['material']


@pytest.mark.parametrize('value',['0','-1','nan','inf','abc'])
def test_reject_invalid_coefficient(tmp_path,value):
    with pytest.raises(ValueError):model.save_coefficient(tmp_path,model.source()['rows'][2]['sku'],value)
    assert not list(tmp_path.rglob('*.json'))


def test_aluminium_not_invented():
    value=model.calculate(model.source()['rows'][0],model.source()['rates'])
    assert value['material'] is None and value['calculated_purchase'] is None


def test_http_csrf_export_filter_and_saving(tmp_path):
    app=Flask(__name__);app.secret_key='test'
    app.config['DETAIL_CALCULATOR_PATHS']={'copy':tmp_path/'copy.json'}
    app.register_blueprint(comparison)
    client=app.test_client()
    assert client.post('/detail-comparison',data={'action':'coefficient'}).status_code==400
    with client.session_transaction() as session:token=session['comparison_csrf']
    sku=model.source()['rows'][2]['sku']
    response=client.post('/detail-comparison',data={'csrf_token':token,'action':'category_coefficient','category':'BACK-SREW','coefficient':'3'})
    assert response.status_code==302
    response=client.get('/detail-comparison',query_string={'q':sku,'export':'csv'})
    data=list(csv.reader(io.StringIO(response.data.decode('utf-8-sig')),delimiter=';'))
    assert len(data)==2 and data[1][0]==sku and data[1][6]=='3,000000'
    assert float(data[1][5].replace(',','.'))==pytest.approx(float(data[1][8].replace(',','.'))*3,abs=2e-6)


def test_full_page_navigation_and_auth(monkeypatch,tmp_path):
    from test_webapp import load_webapp
    monkeypatch.setenv('PRODUCT_ENGINE_SHARED_DATA_DIR',str(tmp_path/'shared'))
    web=load_webapp(monkeypatch,tmp_path)
    client=web.app.test_client()
    response=client.get('/detail-comparison')
    assert response.status_code==200
    text=response.get_data(as_text=True)
    assert '1884 katalogo pozicijos' in text
    assert 'name="coefficient"' in text
    assert '/detail-comparison' in text
    assert client.get('/detail-comparison?page=19').status_code==200
    assert client.get('/detail-comparison?page=bad').status_code==400
    monkeypatch.setattr(web,'auth_enabled',lambda:True)
    assert client.get('/detail-comparison').status_code==302
    assert client.get('/detail-comparison?export=csv').status_code==302
    assert client.get('/detail-comparison?export=xlsx').status_code==302


def test_category_bulk_scope_and_legacy_override(tmp_path):
    _, before = model.results(tmp_path)
    target = next(r for r in before if r['category'] == 'BACK-SREW')
    model.save_coefficient(tmp_path, target['sku'], 9)
    model.save_category_coefficient(tmp_path, 'BACK-SREW', '2,7')
    _, after = model.results(tmp_path)
    changed = [r for r in after if r['category'] == 'BACK-SREW']
    assert len(changed) == 196
    assert {r['sku'].split('-')[0] for r in changed} == {'EU','US'}
    for old, new in zip(before, after):
        if new['category'] == 'BACK-SREW':
            assert new['coefficient'] == 2.7
            assert new['calculated_purchase'] == pytest.approx(new['material'] * 2.7)
            assert new['purchase'] == old['purchase']
        else:
            assert new == old
    for value in ['nan','0','-2']:
        with pytest.raises(ValueError): model.save_category_coefficient(tmp_path, 'PNL', value)
    with pytest.raises(ValueError): model.save_category_coefficient(tmp_path, '../fake', 2)


def test_category_mapping_catalog():
    rows = model.source()['rows']
    groups = {}
    for row in rows:
        groups.setdefault(model.category(row), set()).add(row['coefficient'])
    assert all(len(values) == 1 for values in groups.values())
    assert groups['PNL'] == {2}
    assert groups['PCL'] == {2.2}
    assert groups['LED'] == groups['LEDROD'] == groups['ROD'] == {4}
    assert model.category(dict(rows[2], sku='EUB-C-CAB01-PNL001')) == 'PNL'
    assert model.category(dict(rows[2], sku='USB-C-CAB02-SLF006')) == 'SHELF'


def test_excel_export_current_settings_filters_and_types(tmp_path):
    from openpyxl import load_workbook
    app = Flask(__name__); app.secret_key = 'test'
    app.config['DETAIL_CALCULATOR_PATHS'] = {'copy': tmp_path/'copy.json'}
    app.register_blueprint(comparison)
    model.save_category_coefficient(tmp_path/'furnibox_comparison', 'BACK-SREW', 3.5)
    client = app.test_client()
    response = client.get('/detail-comparison?export=xlsx&page=19')
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    book = load_workbook(io.BytesIO(response.data), data_only=False)
    sheet = book['Detalių kainos']
    assert sheet.max_row == 1889 and sheet.freeze_panes == 'C6'
    assert sheet['A8'].value == model.source()['rows'][2]['sku']
    assert sheet['G8'].value == 3.5
    assert sheet['F8'].value == pytest.approx(sheet['I8'].value * 3.5)
    assert sheet['K8'].number_format == '0.00%'
    assert sheet['I6'].value == 'Netaikoma'
    assert 'Naudoti parametrai' in book.sheetnames
    from webapp.furnibox_comparison import COLUMNS
    _, expected = model.results(tmp_path/'furnibox_comparison')
    for cells, row in zip(sheet.iter_rows(min_row=6), expected):
        for cell, (key, _) in zip(cells, COLUMNS):
            if isinstance(row[key], (int, float)):
                assert cell.value == pytest.approx(row[key])
            assert cell.data_type != 'f'
    response = client.get('/detail-comparison?export=xlsx&q=EU-BACK-SREW-1187x379-BB')
    filtered = load_workbook(io.BytesIO(response.data))['Detalių kainos']
    assert filtered.max_row == 6
    response = client.get('/detail-comparison?export=xlsx&q=nonexistent_xyz')
    assert load_workbook(io.BytesIO(response.data))['Detalių kainos'].max_row == 5
