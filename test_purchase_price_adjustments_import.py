from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from purchase_price_adjustments_import import (
    load_purchase_price_excel_adjustments,
)


class PurchasePriceExcelImportTests(unittest.TestCase):
    def temporary_path(self, directory: str, filename: str) -> Path:
        return Path(directory) / filename

    def remove_worksheet_dimensions(self, path: Path) -> None:
        rewritten = path.with_name(f"rewritten-{path.name}")
        with ZipFile(path, "r") as source, ZipFile(
            rewritten,
            "w",
            ZIP_DEFLATED,
        ) as destination:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet"):
                    payload = re.sub(
                        rb"<dimension\s+ref=\"[^\"]+\"\s*/>",
                        b"",
                        payload,
                    )
                destination.writestr(item, payload)
        rewritten.replace(path)

    def test_imports_main_tamara_adjustments_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.temporary_path(
                directory,
                "Tamara_Adjusted_Component_Prices.xlsx",
            )
            workbook = Workbook()
            workbook.active.title = "INFO"
            sheet = workbook.create_sheet("TAMARA ADJUSTMENTS")
            sheet.append([
                "Internal Reference",
                "Adjusted Purchase Price",
                "Real Purchase Price (reference)",
                "Comment",
            ])
            sheet.append(["7001730", 1.48, 0, "Tamara"])
            sheet.append(["3284903", 0.04, 0, "Tamara"])
            sheet.append(["3503786", 0, 0, "Missing"])
            workbook.save(path)

            result = load_purchase_price_excel_adjustments(path)

            self.assertEqual(result, {
                "7001730": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 1.48,
                },
                "3284903": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 0.04,
                },
            })

    def test_import_does_not_require_worksheet_dimension_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.temporary_path(
                directory,
                "Tamara_Adjusted_Component_Prices.xlsx",
            )
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "TAMARA ADJUSTMENTS"
            sheet.append([
                "Internal Reference",
                "Adjusted Purchase Price",
                "Real Purchase Price (reference)",
            ])
            sheet.append(["3503786", 3.5, 0])
            workbook.save(path)
            self.remove_worksheet_dimensions(path)

            result = load_purchase_price_excel_adjustments(path)

            self.assertEqual(result, {
                "3503786": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 3.5,
                },
            })

    def test_imports_only_changed_prices_from_generated_furnibox_report(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.temporary_path(
                directory,
                "Furnibox_Tamara_Purchase_Prices.xlsx",
            )
            workbook = Workbook()
            components = workbook.active
            components.title = "COMPONENTS"
            headers = [
                "Internal Reference",
                "Name",
                "Price Source",
                "Vendor / Supply Source",
                "Real Furnibox Purchase Price",
                "Furnibox (Tamara) Purchase Price",
                "Reform Markup Factor",
                "Reform Purchase Price",
                "Status / BOM Source",
            ]
            components.append(headers)
            components.append([
                "UNCHANGED", "Part", "LAST PURCHASE PRICE", "Vendor",
                4.25, 4.25, 1, 4.25, "OK",
            ])
            components.append([
                "3503786", "Part", "LAST PURCHASE PRICE", "Vendor",
                0, 3.50, 1, 3.50, "MISSING PRICE",
            ])
            components.append([
                "STILL-MISSING", "Part", "LAST PURCHASE PRICE", "Vendor",
                0, 0, 1, 0, "MISSING PRICE",
            ])
            cabinet_parts = workbook.create_sheet("CABINET PARTS")
            cabinet_parts.append(headers)
            workbook.save(path)

            result = load_purchase_price_excel_adjustments(path)

            self.assertEqual(result, {
                "3503786": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 3.5,
                },
            })

    def test_imports_changed_prices_from_both_generated_report_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.temporary_path(
                directory,
                "Furnibox_Tamara_Purchase_Prices.xlsx",
            )
            workbook = Workbook()
            headers = [
                "Internal Reference",
                "Real Furnibox Purchase Price",
                "Furnibox (Tamara) Purchase Price",
            ]
            components = workbook.active
            components.title = "COMPONENTS"
            components.append(headers)
            components.append(["COMPONENT-1", 0, 1.25])
            cabinet_parts = workbook.create_sheet("CABINET PARTS")
            cabinet_parts.append(headers)
            cabinet_parts.append(["CABINET-PART-1", 2, 2.5])
            workbook.save(path)

            result = load_purchase_price_excel_adjustments(path)

            self.assertEqual(result, {
                "COMPONENT-1": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 1.25,
                },
                "CABINET-PART-1": {
                    "excel_real_price": 2.0,
                    "new_adjustment": 2.5,
                },
            })

    def test_imports_changed_prices_from_legacy_single_report_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.temporary_path(
                directory,
                "Furnibox_Tamara_Purchase_Prices.xlsx",
            )
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "FURNIBOX PURCHASE PRICES"
            sheet.append([
                "Internal Reference",
                "Real Furnibox Purchase Price",
                "Furnibox (Tamara) Purchase Price",
            ])
            sheet.append(["LEGACY-PART", 0, 7.75])
            workbook.save(path)

            result = load_purchase_price_excel_adjustments(path)

            self.assertEqual(result, {
                "LEGACY-PART": {
                    "excel_real_price": 0.0,
                    "new_adjustment": 7.75,
                },
            })


if __name__ == "__main__":
    unittest.main()
