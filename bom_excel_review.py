"""Ticket #28: read-only review core. No Odoo writes or automatic commands.

The baseline must be a separately retained, trusted original export, never the
hidden snapshot supplied by the person returning the edited workbook.
"""
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
import re
from urllib.parse import urlsplit

from openpyxl import load_workbook

IMMUTABLE = ('Produkto kodas', 'Produkto pavadinimas', 'Kategorija',
             'BOM reference', 'BOM tipas', 'TECH BOM write_date')
EDITABLE = ('Eiliskumas', 'Operacija', 'Darbo centras', 'Laikas (min:sek)')
REQUIRED = IMMUTABLE + EDITABLE + ('Routing', 'Pastaba', 'TECH BOM ID',
                                 'TECH operacijos ID', 'TECH operacijos write_date')


def text(value):
    return '' if value is None else str(value).strip()


def integer(value, *, optional=False):
    if optional and not text(value):
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError('Expected an integer') from exc
    if not number.is_finite() or number != number.to_integral_value() or number < 0:
        raise ValueError('Expected a nonnegative integer')
    return int(number)


def minutes(value, *, baseline=False):
    if baseline and isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        result = Decimal(str(value))
        if result.is_finite() and result >= 0:
            return result
        raise ValueError('Invalid baseline minutes')
    if not isinstance(value, str) or not re.fullmatch(r'[0-9]+:[0-5][0-9]', value.strip()):
        raise ValueError('Use minutes:seconds, e.g. 00:30 or 120:00')
    m, s = map(int, value.strip().split(':'))
    return Decimal(m) + Decimal(s) / 60


def require_stage(url):
    parsed = urlsplit(url)
    if (parsed.scheme != 'https' or parsed.hostname != 'stageodoo.furnibox.lt'
            or parsed.port not in (None, 443) or parsed.username or parsed.password
            or parsed.query or parsed.fragment or parsed.path not in ('', '/')):
        raise ValueError('Only https://stageodoo.furnibox.lt is permitted')


def digest(value):
    return sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                             default=str, separators=(',', ':')).encode()).hexdigest()


def read_rows(path):
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        ws = wb['Korekcijos']
        values = iter(ws.iter_rows())
        headers = [text(c.value) for c in next(values)]
        if len(headers) != len(set(headers)) or not set(REQUIRED) <= set(headers):
            raise ValueError('Missing or duplicate Korekcijos headers')
        result = []
        for rownum, cells in enumerate(values, 2):
            if any(c.data_type == 'f' for c in cells):
                raise ValueError(f'Formula in row {rownum}')
            row = dict(zip(headers, (c.value for c in cells)))
            if not any(text(v) for v in row.values()):
                continue
            row['_row'] = rownum
            result.append(row)
        return result
    finally:
        wb.close()


def operation(row, *, baseline=False):
    if not text(row['Operacija']) or not text(row['Darbo centras']):
        raise ValueError('Operation name and work center are required')
    return {'name': text(row['Operacija']), 'sequence': integer(row['Eiliskumas']),
            'workcenter': text(row['Darbo centras']),
            'minutes': str(minutes(row['Laikas (min:sek)'], baseline=baseline))}


def review_rows(baseline, returned):
    errors, warnings, changes, requests = [], [], [], []
    boms, old_ops, seen = {}, {}, set()
    for row in baseline:
        bid = integer(row['TECH BOM ID'])
        oid = integer(row['TECH operacijos ID'], optional=True)
        if bid in boms and any(text(row[k]) != text(boms[bid][k]) for k in IMMUTABLE + ('Routing',)):
            raise ValueError('Trusted baseline has inconsistent BOM metadata')
        boms[bid] = row
        if oid is not None:
            key = bid, oid
            if key in old_ops:
                raise ValueError('Trusted baseline has duplicate operation IDs')
            old_ops[key] = row
    routes = defaultdict(set)
    for row in returned:
        n = row['_row']
        try:
            bid = integer(row['TECH BOM ID'])
            oid = integer(row['TECH operacijos ID'], optional=True)
            if bid not in boms:
                raise ValueError('Unknown source BOM ID')
            base = boms[bid]
            if any(text(row[k]) != text(base[k]) for k in IMMUTABLE):
                raise ValueError('Product/BOM metadata or source date changed')
            routes[bid].add(text(row['Routing']))
            note = text(row['Pastaba'])
            if note:
                requests.append({'row': n, 'source_bom_id': bid,
                                 'source_operation_id': oid, 'note': note,
                                 'requires_human_review': True})
            if oid is None:
                if text(row['TECH operacijos write_date']):
                    raise ValueError('New operation has an existing-operation date')
                if not any(text(row[k]) for k in EDITABLE):
                    continue
                proposed = operation(row)
                changes.append({'kind': 'create', 'row': n, 'source_bom_id': bid,
                                'before': None, 'after': proposed})
            else:
                key = bid, oid
                if key not in old_ops:
                    raise ValueError('Unknown operation or wrong parent BOM')
                if key in seen:
                    raise ValueError('Duplicate existing operation ID')
                seen.add(key)
                old = old_ops[key]
                if text(row['TECH operacijos write_date']) != text(old['TECH operacijos write_date']):
                    raise ValueError('Source operation date changed')
                if any(not text(row[k]) for k in EDITABLE):
                    raise ValueError('Blank existing field is not a deletion request')
                before, after = operation(old, baseline=True), operation(row)
                if before != after:
                    changes.append({'kind': 'update', 'row': n, 'source_bom_id': bid,
                                    'source_operation_id': oid, 'before': before, 'after': after})
        except (ValueError, KeyError, TypeError) as exc:
            errors.append({'row': n, 'error': str(exc)})
    for bid, values in routes.items():
        if len(values) != 1:
            errors.append({'source_bom_id': bid, 'error': 'Conflicting Routing values'})
            continue
        after = next(iter(values))
        before = text(boms[bid]['Routing'])
        if after != before:
            if not after:
                errors.append({'source_bom_id': bid, 'error': 'Routing clear requires explicit review'})
            else:
                changes.append({'kind': 'routing', 'source_bom_id': bid,
                                'before': before, 'after': after})
    for bid, oid in sorted(old_ops.keys() - seen):
        warnings.append({'source_bom_id': bid, 'source_operation_id': oid,
                         'warning': 'Missing row; operation will not be deleted'})
    affected = {c['source_bom_id'] for c in changes}
    for bid in sorted(affected):
        if text(boms[bid]['BOM tipas']) == 'Kit':
            errors.append({'source_bom_id': bid, 'error': 'Kit BOM requires separate review'})
    # Stable for reordered rows, including multiplicity of identical new operations.
    semantic = [{k: v for k, v in c.items() if k != 'row'} for c in changes]
    change_hash = digest(sorted(semantic, key=lambda c: json.dumps(c, sort_keys=True)))
    return {'changes': changes, 'errors': errors, 'warnings': warnings,
            'human_review': requests, 'change_hash': change_hash,
            'proposed_eco_count': len(affected),
            'ready_for_stage_mapping': not errors and not requests,
            'ready_to_write': False}


def unique_mapping(source, candidates):
    """Require a reviewed full identity, never cross-environment numeric IDs.

    Identity includes SKU, company, variant, BOM type, quantity, UOM, reference
    and component fingerprint. Historical reference exceptions must be reviewed
    separately; this function does not silently strip category suffixes.
    """
    fields = ('sku', 'company', 'variant', 'type', 'quantity', 'uom',
              'reference', 'component_fingerprint')
    if any(k not in source for k in fields):
        raise ValueError('Incomplete source identity')
    matches = [c for c in candidates if c.get('active') is True
               and all(k in c and c[k] == source[k] for k in fields)]
    if len(matches) != 1:
        raise ValueError(f'Expected one Stage BOM match; found {len(matches)}')
    return matches[0]


def operation_mapping(source, target):
    """A mismatch or repeated signature blocks rather than guessing by order."""
    fields = ('name', 'sequence', 'workcenter', 'minutes')
    def signature(row):
        return tuple(row[k] for k in fields)
    left, right = Counter(map(signature, source)), Counter(map(signature, target))
    if left != right or any(n != 1 for n in left.values()):
        raise ValueError('Ambiguous or changed Stage operations')
    by_signature = {signature(r): r['id'] for r in target}
    return {r['id']: by_signature[signature(r)] for r in source}
