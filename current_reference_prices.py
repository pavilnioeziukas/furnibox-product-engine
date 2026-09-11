"""User supplied current selling prices, used only for comparison."""
import json
from copy import copy
from functools import lru_cache
from pathlib import Path

from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

COLUMN = "Dabartinė kaina, €"


@lru_cache(maxsize=1)
def prices():
    path = Path(__file__).parent / "manifest" / "current_reference_prices.json"
    return json.loads(path.read_text(encoding="utf-8-sig"))["prices"]


def current_price(sku):
    # No suffix substitution: a different SKU needs its own supplied price.
    return prices().get(str(sku).strip())


def add_to_workbook(workbook):
    for name in ("SO LINE PRICES", "PRICE RESULTS"):
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        headers = [cell.value for cell in sheet[1]]
        if COLUMN in headers:
            continue  # Preserve the price snapshot of an already enriched file.
        sku_column = headers.index("SKU") + 1
        column = headers.index("Final Reform SO Unit Price") + 2
        for index in range(sheet.max_column, column - 1, -1):
            old = get_column_letter(index)
            new = get_column_letter(index + 1)
            sheet.column_dimensions[new].width = sheet.column_dimensions[old].width
        sheet.insert_cols(column)
        cell = sheet.cell(1, column, COLUMN)
        cell._style = copy(sheet.cell(1, column - 1)._style)
        cell.comment = Comment(
            "Dabartinės kainos iš vartotojo 2026-09-11 sąrašo, EUR/vnt. "
            "Tik palyginimui. Tuščia reikšmė: tikslaus SKU kainos nėra arba "
            "šaltinyje N/A / –. SLF-PINS-HRD-4: patvirtinta 0,4213 EUR.",
            "Furnibox",
        )
        sheet.column_dimensions[get_column_letter(column)].width = 23
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column, current_price(sheet.cell(row, sku_column).value))
            cell._style = copy(sheet.cell(row, column - 1)._style)
            cell.number_format = '0.0000 [$€-x-euro2]'
        sheet.auto_filter.ref = sheet.dimensions
