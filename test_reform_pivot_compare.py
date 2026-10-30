import io
import re

import pytest
from openpyxl import Workbook, load_workbook

from test_reform_workspace import setup
from test_reform_approval_workflow import login
from webapp.reform_workspace import db, baseline, unpack
from webapp.reform_pivot_compare import parse, compare, PivotError, displayed_report


def xlsx(rows, title='Product MBOMs - Pivot'):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(['Sum of Part Qty.'])
    ws.append(['SKU', 'MATERIAL_SKU', 'Total'])
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def send(client, path, content, filename='Poul.xlsx', **fields):
    client.get('/reform/pivot-compare')
    with client.session_transaction() as session:
        csrf = session['reform_csrf']
    return client.post('/reform/' + path, data={'csrf': csrf, 'file': (io.BytesIO(content), filename), **fields}, follow_redirects=True)


def seed(client, rows):
    response = send(client, 'pivot-compare/reference', xlsx(rows), label='BOM v10')
    assert response.status_code == 200
    return re.search(r'name="reference_id" value="([^"]+)"', response.text).group(1)


def test_complete_pivot_comparison_without_odoo_or_exchange_token(setup):
    app, first = setup
    second = login(app, 'other', 'other-password')
    with app.app_context(), db() as conn:
        original = baseline(conn)
    rid = seed(first, [('NOT-IN-ODOO', 'Supplier-001', 2), ('NOT-IN-ODOO', 'Old-part', 1), ('REMOVED-PARENT', 12345, 4), ('UNCHANGED', 'A', 1)])
    response = send(second, 'pivot-compare', xlsx([('UNCHANGED', 'A', '1,00'), ('NOT-IN-ODOO', 'Supplier-001', 3), ('NEW-PARENT', 'New-component', 5)]), reference_id=rid, scope='all')
    assert response.status_code == 200 and 'Quantity changed' in response.text
    with app.app_context(), db() as conn:
        row = conn.execute('SELECT * FROM pivot_comparisons').fetchone()
        report = unpack(row['payload'])
        key = row['id']
        assert (report['added'], report['removed'], report['changed'], report['unchanged']) == (1, 2, 1, 1)
        assert row['actor'] == 'other'
        assert baseline(conn) == original
        assert conn.execute('SELECT count(*) FROM drafts').fetchone()[0] == 0
        assert conn.execute('SELECT count(*) FROM submissions').fetchone()[0] == 0
    download = first.get(f'/reform/pivot-compare/results/{key}/download')
    out = load_workbook(io.BytesIO(download.data))
    assert out['Changes'].max_row == 5
    assert [c.value for c in out['Changes'][1]][1:3] == ['SKU', 'MATERIAL_SKU']
    assert out['Sources']['B1'].value == 'BOM v10'


def test_reference_replacement_does_not_rewrite_existing_report(setup):
    app, client = setup
    rid = seed(client, [('P', 'C', 1)])
    response = send(client, 'pivot-compare', xlsx([('P', 'C', 2)]), reference_id=rid)
    url = response.request.path
    before = client.get(url).data
    new_rid = seed(client, [('P', 'C', 100)])
    assert new_rid != rid
    assert client.get(url).data == before
    with app.app_context(), db() as conn:
        assert conn.execute('SELECT count(*) FROM pivot_references').fetchone()[0] == 2


def test_row_order_numeric_codes_and_zero_are_compared_as_values():
    a = parse(xlsx([('SKU', 12345, 1), ('SKU', '00012', 2)]))
    b = parse(xlsx([('SKU', '00012', '2.0'), ('SKU', 12345, 0)]))
    result = compare(a, b)
    assert result['unchanged'] == 1 and result['changed'] == 1
    assert result['changes'][0]['difference'] == '-1'
    assert result['changes'][0]['status'] == 'Quantity changed'


def test_all_validation_issues_are_reported_and_nothing_is_saved(setup):
    app, client = setup
    rid = seed(client, [('P', 'C', 1)])
    response = send(client, 'pivot-compare', xlsx([('P', 'C', 1), ('P', 'C', 2), ('P', None, 2), ('Q', 'F', '=1+1'), ('Z', 'F', 'NaN')]), reference_id=rid)
    assert response.status_code == 400
    assert '4 issues' in response.text and 'Duplicate' in response.text and 'formulas' in response.text
    with app.app_context(), db() as conn:
        assert conn.execute('SELECT count(*) FROM pivot_comparisons').fetchone()[0] == 0


def test_comparison_works_without_any_odoo_catalogue_or_open_version(setup):
    app, client = setup
    with app.app_context(), db() as conn:
        conn.execute('DELETE FROM state')
    rid = seed(client, [('P', 'C', 1)])
    page = send(client, 'pivot-compare', xlsx([('P', 'C', 1)]), reference_id=rid)
    assert 'No differences' in page.text


def test_unknown_reference_and_anonymous_access_are_rejected(setup):
    app, client = setup
    assert app.test_client().get('/reform/pivot-compare').status_code == 302
    response = send(client, 'pivot-compare', xlsx([('P', 'C', 1)]), reference_id='missing')
    assert response.status_code == 409
    template = client.get('/reform/pivot-compare/template')
    wb = load_workbook(io.BytesIO(template.data))
    assert wb.sheetnames[0] == 'Product MBOM Pivot' and '_Exchange' not in wb


def test_template_roundtrip_and_exact_code_changes():
    old = parse(xlsx([('P', 'part', 1), ('P', 'Zero', 0)]))
    new = parse(xlsx([('P', 'PART', 1), ('P', 'Zero', 0)], title='Product MBOM Pivot'))
    report = compare(old, new, 'all')
    assert (report['added'], report['removed'], report['changed'], report['unchanged']) == (1, 1, 0, 1)


def test_excel_macro_workbook_extension_and_invalid_file_message(setup):
    app, client = setup
    rid = seed(client, [('P', 'C', 1)])
    assert 'No differences' in send(client, 'pivot-compare', xlsx([('P', 'C', 1)]), filename='Poul.xlsm', reference_id=rid).text
    page = send(client, 'pivot-compare', b'not an Excel file', reference_id=rid)
    assert page.status_code == 400 and 'readable Excel workbook' in page.text


def test_partial_upload_ignores_absent_rows_even_for_same_parent(setup):
    app, client = setup
    rid = seed(client, [('P', 'A', 1), ('P', 'B', 2), ('Q', 'C', 3)])
    response = send(client, 'pivot-compare', xlsx([('P', 'A', 4)]), reference_id=rid)
    url = response.request.path
    with app.app_context(), db() as conn:
        item = conn.execute('SELECT * FROM pivot_comparisons').fetchone()
        report = unpack(item['payload'])
        assert (report['changed'], report['removed'], report['ignored']) == (1, 0, 2)
    result_id = url.rsplit('/', 1)[1]
    out = load_workbook(io.BytesIO(client.get(f'{url}/download').data))
    assert out['Changes'].max_row == 2
    assert 'Complete version; missing rows are removals' in client.get(url + '?scope=all').text
    full = load_workbook(io.BytesIO(client.get(f'{url}/download?scope=all').data))
    assert full['Changes'].max_row == 4


def test_legacy_full_report_reconstructs_uploaded_rows_without_reupload(setup):
    app, client = setup
    original = parse(xlsx([('P', 'A', 1), ('P', 'B', 2), ('Q', 'C', 3), ('Q', 'D', 4)]))
    uploaded = parse(xlsx([('P', 'A', 5), ('Q', 'C', 3)]))
    legacy = compare(original, uploaded, 'all')
    legacy.pop('scope')
    assert legacy['removed'] == 2
    assert displayed_report(original, legacy)['changed'] == 1
    assert displayed_report(original, legacy)['removed'] == 0
    assert displayed_report(original, legacy)['unchanged'] == 1
    rid = seed(client, [('P', 'A', 1), ('P', 'B', 2), ('Q', 'C', 3), ('Q', 'D', 4)])
    from webapp.reform_workspace import pack
    with app.app_context(), db() as conn:
        conn.execute('INSERT INTO pivot_comparisons VALUES (?,?,?,?,?,?,?)', ('legacy', rid, 'Old-upload.xlsx', 'paul', '2026-10-30', 'hash', pack(legacy)))
    page = client.get('/reform/pivot-compare/results/legacy')
    assert 'Only uploaded rows; missing rows are unchanged' in page.text
    out = load_workbook(io.BytesIO(client.get('/reform/pivot-compare/results/legacy/download').data))
    assert out['Changes'].max_row == 2
    with app.app_context(), db() as conn:
        assert unpack(conn.execute("SELECT payload FROM pivot_comparisons WHERE id='legacy'").fetchone()[0]) == legacy


def test_partial_26_row_upload_against_7333_row_reference():
    before = [{'sku': 'SKU', 'material': str(i), 'quantity': '1', 'row': i + 3} for i in range(7333)]
    after = [{**r, 'quantity': '2' if i < 3 else '1'} for i, r in enumerate(before[:26])]
    old = compare(before, after, 'all')
    old.pop('scope')
    assert (old['removed'], old['changed'], old['unchanged']) == (7307, 3, 23)
    fixed = displayed_report(before, old)
    assert (fixed['removed'], fixed['changed'], fixed['unchanged'], len(fixed['changes'])) == (0, 3, 23, 3)
