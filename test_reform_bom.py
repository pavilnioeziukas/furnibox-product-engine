import copy
import io

import pytest
from openpyxl import load_workbook
from test_reform_workspace import setup, sample
from test_reform_excel import upload_book
from webapp.reform_workspace import db, baseline, pack, unpack


def export(client, query='scope=all'):
    response = client.get('/reform/excel/download?' + query)
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.data))


def saved(app):
    with app.app_context(), db() as conn:
        return unpack(conn.execute('SELECT payload FROM drafts').fetchone()[0])['target']


def test_default_vertical_roundtrip_and_supplier_reference(setup):
    app, client = setup
    with app.app_context(), db() as conn:
        conn.execute('INSERT INTO state VALUES (3,?)', (pack({'PANEL-01': ['SUP-1']}),))
    b = export(client)
    assert b.sheetnames[:2] == ['BOM', 'Non-BOM'] and 'BOM Map' not in b
    assert b['Components'].sheet_state == 'veryHidden'
    assert b['BOM']['E2'].value == 'SUP-1'
    assert b['BOM']['B2'].font.sz == 9 and b['BOM'].freeze_panes == 'B2'
    assert b['BOM'].column_dimensions['H'].hidden and b['BOM'].column_dimensions['I'].hidden
    assert 'No changes detected' in upload_book(client, b).text
    b['BOM']['E2'] = 'EDITED'
    assert 'reference field' in upload_book(client, b).text
    b['BOM']['E2'] = 'SUP-1'
    b['BOM']['C2'] = 7
    assert 'Excel changes saved' in upload_book(client, b).text
    assert saved(app)['boms']['1']['components'][0]['quantity'] == 7
    assert 'Your draft has changed' in upload_book(client, b).text


def test_direct_children_shared_sub_bom_once_and_output_quantity(setup):
    app, client = setup
    data = sample()
    data['products']['CAB-02'] = dict(data['products']['CAB-01'], sku='CAB-02')
    data['boms']['2'] = dict(copy.deepcopy(data['boms']['1']), id='2', sku='CAB-02')
    data['boms']['3'] = dict(copy.deepcopy(data['boms']['1']), id='3', sku='PANEL-01', quantity=2,
                           components=[{'id': '31', 'sku': 'HINGE-01', 'quantity': 6, 'uom': 'vnt.', 'uom_id': 1}])
    with app.app_context(), db() as conn:
        conn.execute('UPDATE state SET payload=? WHERE id=1', (pack(data),))
    b = export(client)
    rows = list(b['BOM'].values)[1:]
    assert len(rows) == 5  # Two cabinets with two direct children, one shared sub-BOM row.
    assert [(r[0], r[1], r[2]) for r in rows if r[0] == 'PANEL-01'] == [('PANEL-01', 'HINGE-01', 6)]
    assert 'No changes detected' in upload_book(client, b).text
    b['BOM']['C6'] = 8
    assert 'Excel changes saved' in upload_book(client, b).text
    target = saved(app)
    assert target['boms']['3']['components'][0]['quantity'] == 8
    assert target['boms']['1'] == data['boms']['1'] and target['boms']['2'] == data['boms']['2']


def test_removal_addition_and_row_reordering(setup):
    app, client = setup
    b = export(client)
    rows = list(b['BOM'].values)[1:]
    b['BOM'].delete_rows(2, 2)
    for row in reversed(rows):
        b['BOM'].append(row)
    assert 'No changes detected' in upload_book(client, b).text
    b['BOM']['F2'] = 'REMOVE'
    b['BOM'].append(['CAB-01', 'HINGE-01', 9, None, None, 'KEEP'])
    assert 'Excel changes saved' in upload_book(client, b).text
    assert [(r['sku'], r['quantity']) for r in saved(app)['boms']['1']['components']] == [('PANEL-01', 2), ('HINGE-01', 9)]


def test_blank_template_new_product_bom_and_non_bom_cards(setup):
    app, client = setup
    data = sample()
    for p in data['products'].values():
        p.update(category='Parts', category_id=7)
    with app.app_context(), db() as conn:
        conn.execute('UPDATE state SET payload=? WHERE id=1', (pack(data),))
    b = export(client, '')
    assert b['BOM'].max_row == 1
    b['Products'].append(['NEW-CAB', 'New cabinet', 'Parts', 'vnt.', None])
    b['Non-BOM'].append(['NEW-PART', 'New component', 'Parts', 'vnt.', None])
    b['BOMs'].append(['NEW-1', 'NEW-CAB', 'New BOM', 2])
    b['BOM'].append(['NEW-CAB', 'NEW-PART', 3, None, None, 'KEEP'])
    assert 'Excel changes saved' in upload_book(client, b).text
    new = next(v for v in saved(app)['boms'].values() if v['sku'] == 'NEW-CAB')
    assert new['quantity'] == 2 and new['components'][0]['quantity'] == 3


@pytest.mark.parametrize('cell,value,error', [
    ('A2', 'OTHER', 'reference field'), ('I2', 'fake', 'Line ID is invalid'),
    ('C2', 0, 'greater than zero'), ('C2', '=1+1', 'replace formulas'),
    ('B2', 'APACK-PRIVATE', 'component does not exist'), ('B2', 'CAB-01', 'cycle'),
])
def test_invalid_uploads_are_atomic(setup, cell, value, error):
    app, client = setup
    b = export(client)
    b['BOM'][cell] = value
    assert error in upload_book(client, b).text
    with app.app_context(), db() as conn:
        assert conn.execute('SELECT count(*) FROM drafts').fetchone()[0] == 0


def test_missing_duplicate_rows_and_read_only_bom(setup):
    app, client = setup
    b = export(client)
    b['BOM'].delete_rows(2)
    assert 'exported rows are missing' in upload_book(client, b).text
    b = export(client)
    b['BOM'].append([c.value for c in b['BOM'][2]])
    assert 'Line ID is invalid or duplicated' in upload_book(client, b).text
    with app.app_context(), db() as conn:
        data = baseline(conn)
        data['boms']['1']['read_only'] = 'Internal assembly'
        conn.execute('UPDATE state SET payload=? WHERE id=1', (pack(data),))
    b = export(client)
    assert 'No changes detected' in upload_book(client, b).text
    b['BOM']['C2'] = 9
    assert 'read-only' in upload_book(client, b).text
