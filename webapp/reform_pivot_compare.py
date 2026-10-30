"""Independent, immutable Pivot-to-Pivot comparisons. No Odoo or proposal writes."""
import hashlib
import io
import math
import secrets
from decimal import Decimal, InvalidOperation

from flask import abort, flash, redirect, render_template, request, send_file, url_for, has_request_context
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from webapp.reform_workspace import reform, db, owner, now, pack, unpack
from webapp.reform_excel import LIMIT, read_book
from webapp.reform_pivot import ALIASES, NAME, HEADERS


class PivotError(ValueError):
    def __init__(self, issues):
        self.issues = issues
        super().__init__('Pivot validation failed.')


def code(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse(content):
    try:
        wb = read_book(content)
    except ValueError as exc:
        raise PivotError([{'row': '', 'sku': '', 'material': '', 'reason': 'Choose a readable Excel workbook smaller than 5 MB containing the Pivot sheet.'}]) from exc
    try:
        candidates, issues = [], []
        for name in ALIASES:
            if name not in wb:
                continue
            ws = wb[name]
            if ws.max_row > 20032 or ws.max_column > 20:
                raise PivotError([{'row': '', 'sku': '', 'material': '', 'reason': 'Pivot sheet exceeds the 20,000-row / 20-column limit.'}])
            rows = list(ws.iter_rows())
            header = next((i for i, r in enumerate(rows[:30]) if [code(c.value) for c in r[:3]] == HEADERS), None)
            if header is None:
                issues.append({'row': '', 'sku': '', 'material': '', 'reason': f'{name}: SKU / MATERIAL_SKU / Total headers are required in columns A:C.'})
                continue
            data, seen = [], set()
            for n, cells in enumerate(rows[header + 1:], header + 2):
                values = [c.value for c in cells]
                if not any(code(v) for v in values):
                    continue
                parent, material = code(values[0]), code(values[1])
                reason = ''
                if any(c.data_type == 'f' for c in cells):
                    reason = 'Paste values instead of formulas.'
                elif not parent or not material:
                    reason = 'SKU and MATERIAL_SKU must be present on every data row. Remove subtotals and grand totals.'
                elif any(code(v) for v in values[3:]):
                    reason = 'Only three data columns are supported: SKU / MATERIAL_SKU / Total.'
                elif (parent, material) in seen:
                    reason = 'Duplicate SKU / MATERIAL_SKU pair. Supply one quantity per pair.'
                try:
                    quantity = Decimal(code(values[2]).replace(',', '.'))
                    if (isinstance(values[2], bool) or not quantity.is_finite() or quantity < 0 or
                            not math.isfinite(float(quantity)) or (quantity != 0 and float(quantity) == 0)):
                        raise InvalidOperation()
                except (InvalidOperation, ValueError):
                    quantity = None
                    if not reason:
                        reason = 'Total must be a finite number greater than or equal to zero.'
                if reason:
                    issues.append({'row': n, 'sku': parent, 'material': material, 'reason': reason})
                else:
                    seen.add((parent, material))
                    data.append({'sku': parent, 'material': material, 'quantity': str(quantity), 'row': n})
            if data:
                candidates.append(data)
        if issues:
            raise PivotError(issues)
        if len(candidates) != 1:
            raise PivotError([{'row': '', 'sku': '', 'material': '', 'reason': 'Supply exactly one populated Product MBOM Pivot or Product MBOMs - Pivot sheet.'}])
        return candidates[0]
    finally:
        wb.close()


def compare(before, after, scope='uploaded'):
    new = {(r['sku'], r['material']): r for r in after}
    old = {(r['sku'], r['material']): r for r in before
           if scope == 'all' or (r['sku'], r['material']) in new}
    result, unchanged = [], 0
    for key in sorted(set(old) | set(new)):
        a, b = old.get(key), new.get(key)
        x, y = (Decimal(a['quantity']) if a else None), (Decimal(b['quantity']) if b else None)
        if a and b and x == y:
            unchanged += 1
            continue
        result.append({'sku': key[0], 'material': key[1], 'old': str(x) if a else None,
                       'new': str(y) if b else None, 'difference': str((y or Decimal(0)) - (x or Decimal(0))),
                       'status': 'Added' if not a else 'Removed' if not b else 'Quantity changed',
                       'old_row': a['row'] if a else None, 'new_row': b['row'] if b else None})
    return {'changes': result, 'unchanged': unchanged, 'old_count': len(old), 'new_count': len(after),
            'scope': scope, 'ignored': len(before) - len(old),
            'added': sum(r['status'] == 'Added' for r in result),
            'removed': sum(r['status'] == 'Removed' for r in result),
            'changed': sum(r['status'] == 'Quantity changed' for r in result)}


def displayed_report(before, saved, scope=None):
    """Reconstruct legacy full comparisons without changing their stored evidence."""
    after = saved.get('uploaded_rows')
    if after is None:
        changed = {(r['sku'], r['material']): r for r in saved['changes']}
        after = [r for r in before if (r['sku'], r['material']) not in changed]
        after += [{'sku': r['sku'], 'material': r['material'], 'quantity': r['new'], 'row': r['new_row']}
                  for r in saved['changes'] if r['new'] is not None]
    selected = scope if scope in ('uploaded', 'all') else saved.get('scope', 'uploaded')
    return compare(before, after, selected)


def schema(conn):
    conn.execute('CREATE TABLE IF NOT EXISTS pivot_references (id TEXT PRIMARY KEY, label TEXT, filename TEXT, actor TEXT, created TEXT, sha256 TEXT, payload BLOB)')
    conn.execute('CREATE TABLE IF NOT EXISTS pivot_comparisons (id TEXT PRIMARY KEY, reference_id TEXT, filename TEXT, actor TEXT, created TEXT, sha256 TEXT, payload BLOB)')


def save_reference(conn, rows, label, filename, sha256, reference_id=None):
    schema(conn)
    key = reference_id or secrets.token_hex(16)
    existing = conn.execute('SELECT * FROM pivot_references WHERE id=?', (key,)).fetchone()
    if existing:
        if existing['sha256'] != sha256:
            raise ValueError('An immutable reference with this identifier already exists.')
        return key
    conn.execute('INSERT INTO pivot_references VALUES (?,?,?,?,?,?,?)',
                 (key, label, filename, owner() if has_request_context() else 'Furnibox import', now(), sha256, pack(rows)))
    return key


def reference(conn, key=None):
    schema(conn)
    return (conn.execute('SELECT * FROM pivot_references WHERE id=?', (key,)).fetchone() if key else
            conn.execute('SELECT * FROM pivot_references ORDER BY created DESC LIMIT 1').fetchone())


def upload():
    file = request.files.get('file')
    if not file or not file.filename.lower().endswith(('.xlsx', '.xlsm')):
        raise PivotError([{'row': '', 'sku': '', 'material': '', 'reason': 'Select an .xlsx or .xlsm workbook containing the Pivot sheet.'}])
    raw = file.stream.read(LIMIT + 1)
    return file.filename[:200], raw, parse(raw)


def page(result_id=None, issues=None, status=200):
    with db() as conn:
        schema(conn)
        item = conn.execute('SELECT * FROM pivot_comparisons WHERE id=?', (result_id,)).fetchone() if result_id else None
        if result_id and not item:
            abort(404)
        ref = reference(conn, item['reference_id'] if item else None)
        summary = displayed_report(unpack(ref['payload']), unpack(item['payload']), request.args.get('scope')) if item else None
        query = request.args.get('q', '').strip()[:200]
        changed = summary['changes'] if summary else []
        if query:
            changed = [r for r in changed if query.casefold() in (r['sku'] + ' ' + r['material']).casefold()]
        try:
            number = max(1, int(request.args.get('page', 1)))
        except ValueError:
            number = 1
        pages = max(1, (len(changed) + 99) // 100)
        number = min(number, pages)
        visible = changed[(number - 1) * 100:number * 100]
        base_count = len(unpack(ref['payload'])) if ref else 0
    return render_template('reform_pivot_compare.html', reference=ref, base_count=base_count,
                           result=item, summary=summary, rows=visible, issues=(issues or [])[:200],
                           issue_count=len(issues or []), query=query, page=number, pages=pages), status


@reform.get('/pivot-compare')
def pivot_compare():
    return page()


@reform.post('/pivot-compare/reference')
def pivot_reference_upload():
    try:
        filename, raw, rows = upload()
        label = request.form.get('label', '').strip()[:100]
        if not label:
            raise ValueError('Enter a reference version name, for example BOM v10.')
        with db() as conn:
            save_reference(conn, rows, label, filename, hashlib.sha256(raw).hexdigest())
        flash('Reference Pivot saved. Earlier references and reports are preserved.')
        return redirect(url_for('reform.pivot_compare'))
    except PivotError as exc:
        return page(issues=exc.issues, status=400)
    except ValueError as exc:
        return page(issues=[{'row': '', 'sku': '', 'material': '', 'reason': str(exc)}], status=400)


@reform.post('/pivot-compare')
def pivot_compare_upload():
    try:
        filename, raw, rows = upload()
        with db() as conn:
            ref = reference(conn, request.form.get('reference_id', ''))
            if not ref:
                abort(409, 'Choose an existing reference Pivot before uploading the new version.')
            selected = 'all' if request.form.get('scope') == 'all' else 'uploaded'
            report = compare(unpack(ref['payload']), rows, selected)
            report['uploaded_rows'] = rows
            key = secrets.token_hex(16)
            conn.execute('INSERT INTO pivot_comparisons VALUES (?,?,?,?,?,?,?)',
                         (key, ref['id'], filename, owner(), now(), hashlib.sha256(raw).hexdigest(), pack(report)))
        return redirect(url_for('reform.pivot_compare_result', result_id=key))
    except PivotError as exc:
        return page(issues=exc.issues, status=400)
    except ValueError as exc:
        return page(issues=[{'row': '', 'sku': '', 'material': '', 'reason': str(exc)}], status=400)


@reform.get('/pivot-compare/results/<result_id>')
def pivot_compare_result(result_id):
    return page(result_id)


def workbook(headers, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for values in rows:
        ws.append(values)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='24496B')
        cell.font = Font(bold=True, color='FFFFFF')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 38 if col[0].column in (2, 3) else 22
        for cell in col[1:]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '0.########'
    return wb


@reform.get('/pivot-compare/template')
def pivot_compare_template():
    wb = workbook(HEADERS, [], NAME)
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return send_file(data, as_attachment=True, download_name='Reform-Pivot-Comparison-Template.xlsx')


@reform.get('/pivot-compare/results/<result_id>/download')
def pivot_compare_download(result_id):
    with db() as conn:
        schema(conn)
        item = conn.execute('SELECT * FROM pivot_comparisons WHERE id=?', (result_id,)).fetchone()
        if not item:
            abort(404)
        ref = reference(conn, item['reference_id'])
        report = displayed_report(unpack(ref['payload']), unpack(item['payload']), request.args.get('scope'))
    numeric = lambda q: None if q is None else float(Decimal(q))
    wb = workbook(['Change', 'SKU', 'MATERIAL_SKU', 'Reference quantity', 'New quantity', 'Difference', 'Reference row', 'New row'],
                  [[r['status'], r['sku'], r['material'], numeric(r['old']), numeric(r['new']), numeric(r['difference']), r['old_row'], r['new_row']] for r in report['changes']], 'Changes')
    info = wb.create_sheet('Sources')
    for row in [['Reference version', ref['label']], ['Reference file', ref['filename']], ['Reference SHA256', ref['sha256']],
                ['New version file', item['filename']], ['New version SHA256', item['sha256']], ['Uploaded by', item['actor']],
                ['Compared at (UTC)', item['created']], ['Added', report['added']], ['Removed', report['removed']],
                ['Quantity changed', report['changed']], ['Unchanged', report['unchanged']],
                ['Comparison scope', 'Complete version: absent rows are removals' if report['scope'] == 'all' else 'Uploaded rows only: absent rows are unchanged'],
                ['Reference rows outside scope', report['ignored']],
                ['Comparison', 'No Odoo data used. Codes are compared exactly after trimming surrounding spaces.']]:
        info.append(row)
        for cell in info[info.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    info.column_dimensions['A'].width, info.column_dimensions['B'].width = 28, 110
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return send_file(data, as_attachment=True, download_name='Reform-Pivot-Changes.xlsx')
