"""Editable BOM hierarchy projection with stable references to actual BOM lines."""
import copy
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MAX_ROWS = 100000


def paths(scope):
    by_sku = {b['sku']: (key, b) for key, b in scope['boms'].items()}
    result = []
    def walk(top, key, bom, trail, visited):
        if len(trail) > 20:
            raise ValueError('BOM depth exceeds 20 levels. Furnibox must review this catalogue.')
        for line in bom['components']:
            edge = {'key': key, 'line': line, 'output': bom['quantity'], 'output_uom': bom.get('uom_id'), 'parent': bom['sku']}
            chain = trail + [edge]
            child = by_sku.get(line['sku'])
            if child and child[0] not in visited and child[1]['components']:
                walk(top, child[0], child[1], chain, visited | {child[0]})
            else:
                result.append({'top': top, 'path': chain, 'cycle': bool(child and child[0] in visited)})
                if len(result) > MAX_ROWS:
                    raise ValueError('The hierarchy exceeds 100,000 rows. Export a smaller product selection.')
    for key, bom in sorted(scope['boms'].items(), key=lambda item: item[1]['sku']):
        walk(bom['sku'], key, bom, [], {key})
    return result


def projection(record):
    rows = paths(record['scope'])
    levels = max([1] + [len(row['path']) - 1 for row in rows])
    names = {2:'II',3:'III',4:'IV',5:'V',6:'VI'}
    headers = ['Top BOM SKU']
    for depth in range(levels):
        level = names.get(depth + 2, str(depth + 2))
        headers += [f'Level {level} SKU', f'Level {level} Qty']
    headers += ['Purchased Component SKU', ('Component Qty in Level II' if levels == 1 else 'Component Qty in Final Level'), 'Total Qty in Top BOM',
                'Supplier Code', 'Unit', 'Parent BOM SKU', 'Action', 'Note', 'Row ID']
    rendered = []
    products = record['scope']['products']
    for number, entry in enumerate(rows, 1):
        chain = entry['path']; leaf = chain[-1]['line']; row = [entry['top']]
        intermediate = chain[:-1] if len(chain) > 1 else chain
        for depth in range(levels):
            row += [intermediate[depth]['line']['sku'], intermediate[depth]['line']['quantity']] if depth < len(intermediate) else ['', '']
        total = 1.0; note = []
        for edge in chain:
            total *= edge['line']['quantity'] / edge['output']
            product = products[edge['line']['sku']]
            if edge['line'].get('uom_id') != product.get('uom_id') or edge['output_uom'] != products[edge['parent']].get('uom_id'):
                note.append('Unit conversion required; total not calculated')
        if entry['cycle']:
            note.append('Cyclic BOM: review required')
        if any(record['scope']['boms'][edge['key']].get('read_only') for edge in chain):
            note.append('Includes a read-only BOM')
        codes = record.get('supplier_codes', {}).get(leaf['sku'], products[leaf['sku']].get('supplier_codes', []))
        row += [leaf['sku'], leaf['quantity'] if len(chain) > 1 else 1,
                '' if entry['cycle'] or any('conversion' in n for n in note) else total, '; '.join(codes), leaf['uom'], chain[-1]['parent'], 'KEEP', '; '.join(dict.fromkeys(note)), f'R{number}']
        rendered.append((row, entry))
    return headers, rendered, levels


def style_map(sheet, levels):
    """Separate exported BOM groups without adding rows or changing exchange data."""
    leaf = 2 + 2 * levels
    group_starts = {1, leaf, leaf + 2, leaf + 3, leaf + 6}
    group_starts.update(range(2, leaf, 2))
    divider = Side(style='thin', color='B8C8D8')
    top_rule = Side(style='medium', color='344F6A')
    sub_rule = Side(style='thin', color='A7B9CA')
    no_rule = Side()
    bands = [PatternFill('solid', fgColor=c) for c in ('E4EDF7', 'FCECDD')]
    top_fills = [PatternFill('solid', fgColor=c) for c in ('CFDFEF', 'F3DCC6')]
    header_colors = ['315B86', '456D98', '596C91']
    sheet.sheet_properties.tabColor = '24496B'
    sheet.row_dimensions[1].height = 54
    sheet.print_title_rows = '1:1'
    for cell in sheet[1]:
        col = cell.column
        color = ('24496B' if col == 1 else header_colors[((col - 2) // 2) % 3]
                 if col < leaf else '276B66' if col < leaf + 2 else
                 '3D586E' if col == leaf + 2 else '80613E' if col == leaf + 6 else '586779')
        cell.fill = PatternFill('solid', fgColor=color)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(right=Side(style='thin', color='FFFFFF'))
        title = str(cell.value)
        width = (40 if col == 1 or 'SKU' in title else 18 if 'Qty' in title else
                 12 if title in ('Unit', 'Action') else 55 if title == 'Note' else 30)
        sheet.column_dimensions[cell.column_letter].width = width
    previous_top = previous_parent = None
    stripe = 0
    borders = {}
    body_font = Font(name='Calibri', size=11, color='243746')
    key_font = Font(name='Calibri', size=11, bold=True, color='24496B')
    for row in sheet.iter_rows(min_row=2):
        top, parent = row[0].value, row[leaf + 4].value
        new_top = top != previous_top
        boundary = 2 if new_top else 1 if parent != previous_parent else 0
        if new_top:
            stripe = 0
        sheet.row_dimensions[row[0].row].height = 24
        for cell in row:
            start = cell.column in group_starts
            key = (boundary, start)
            if key not in borders:
                borders[key] = Border(top=top_rule if boundary == 2 else sub_rule if boundary else no_rule,
                                      left=divider if start else no_rule)
            cell.border = borders[key]
            cell.fill = top_fills[stripe % 2] if cell.column == 1 else bands[stripe % 2]
            cell.font = key_font if cell.column == 1 else body_font
            cell.alignment = Alignment(vertical='center', horizontal='right' if isinstance(cell.value, (int, float)) else 'left')
            if isinstance(cell.value, (int, float)):
                cell.number_format = 'General'
        previous_top, previous_parent = top, parent
        stripe += 1


def add_map(wb, record):
    headers, rows, levels = projection(record)
    sheet = wb.create_sheet('BOM Map', 0)
    sheet.append(headers)
    for row, _ in rows:
        sheet.append(row)
    sheet.column_dimensions[get_column_letter(len(headers))].hidden = True
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    non = wb.create_sheet('Non-BOM', 1)
    non.append(['Product code', 'Name', 'Category', 'Unit', 'Supplier Code'])
    parents = {b['sku'] for b in record['scope']['boms'].values()}
    product_sheet = wb['Products']
    product_sheet.cell(1, 5, 'Supplier Code')
    for row in range(product_sheet.max_row, 1, -1):
        code = product_sheet.cell(row, 1).value
        codes = record.get('supplier_codes', {}).get(code, record['scope']['products'][code].get('supplier_codes', []))
        product_sheet.cell(row, 5, '; '.join(codes))
        if code not in parents:
            non.append([product_sheet.cell(row, col).value for col in range(1, 6)])
            product_sheet.delete_rows(row)
    wb['Components'].sheet_state = 'veryHidden'
    for ws in (sheet, non, product_sheet):
        ws.freeze_panes = 'B2' if ws == sheet else 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='4472C4')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 44
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 34 if 'Qty' not in str(ws.cell(1,col).value) else 18
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Calibri', size=11)
                cell.alignment = Alignment(vertical='top')
                cell.fill = PatternFill('solid', fgColor='EAF0F7' if cell.row % 2 == 0 else 'FFFFFF')
    style_map(sheet, levels)
    guide = wb['Instructions']
    guide['B4'] = 'Products: BOM product cards. Non-BOM: cards without a BOM. Edit names/categories or add new cards.'
    guide['B6'] = 'BOM Map: edit quantities or component codes. Use Action=REMOVE for the final component. Do not delete rows.'
    guide['B7'] = 'Keep map rows and hidden Row IDs. Repeated occurrences identify the same underlying BOM line.'
    from openpyxl.worksheet.datavalidation import DataValidation
    for name, col in [('Categories','C'),('Units','D')]:
        if name in wb.defined_names:
            rule = DataValidation(type='list', formula1=name, allow_blank=False)
            rule.showErrorMessage = True
            non.add_data_validation(rule); rule.add(f'{col}2:{col}10001')
    rule = DataValidation(type='list', formula1='"KEEP,REMOVE"', allow_blank=False)
    rule.showErrorMessage = True
    sheet.add_data_validation(rule)
    rule.add(f'{get_column_letter(1+2*levels+7)}2:{get_column_letter(1+2*levels+7)}{MAX_ROWS+1}')
    guide.append(['Main editing sheet', 'Edit BOM Map. Quantities at each level are per parent BOM output. Total Qty is per one Top BOM unit and is recalculated on the next export.'])
    guide.append(['Repeated lines', 'Changing one occurrence updates that same BOM line wherever it is used. Different proposed values for the same line are rejected.'])
    guide.append(['Remove / add', 'Set Action=REMOVE to remove the final component from its parent BOM. Add a row with blank Row ID, Parent BOM SKU, Purchased Component SKU and Component Qty to add a component.'])
    guide.append(['Cards and suppliers', 'Products contains BOM product cards; Non-BOM contains cards without a BOM. Both can be edited. Supplier Code is an Odoo reference field; multiple supplier codes are separated by semicolons.'])
    guide.append(['New BOM', 'Add the product card, define its new BOM in BOMs, then add component rows in BOM Map using Parent BOM SKU.'])
    guide.append(['Reading the map', 'Dark horizontal lines start a Top BOM. Thin lines separate parent BOMs. Blue and peach rows help follow components. Header colours separate hierarchy levels, components and reference fields.'])


def translate(wb, record):
    """Translate map proposals into canonical legacy sheets for the existing validator."""
    from webapp.reform_excel import HEADERS, text, positive
    headers, original, levels = projection(record)
    if 'BOM Map' not in wb or wb['BOM Map'].max_row > MAX_ROWS + 1:
        raise ValueError('BOM Map is missing or exceeds 100,000 rows.')
    iterator = wb['BOM Map'].iter_rows()
    if [c.value for c in next(iterator)] != headers:
        raise ValueError('BOM Map headers changed. Keep the original columns.')
    ref = {row[-1]:(row, entry) for row, entry in original}
    seen, proposals, additions = set(), {}, []
    leafcol = 1 + 2 * levels
    def propose(edge, field, value):
        original_value = edge['line'][field]
        if value == original_value:
            return
        key = (edge['key'], edge['line']['id'], field)
        if key in proposals and proposals[key] != value:
            raise ValueError('The same BOM line has conflicting edits in different map rows. Use one proposed value.')
        proposals[key] = value
    for cells in iterator:
        values = [c.value for c in cells]
        if not any(v is not None for v in values):
            continue
        if len(values) != len(headers) or any(c.data_type == 'f' for c in cells):
            raise ValueError('BOM Map: keep all columns and use values, not formulas.')
        rid = text(values[-1]); action = text(values[leafcol+6])
        if action not in ('KEEP', 'REMOVE'):
            raise ValueError('BOM Map: use Action KEEP or REMOVE.')
        if not rid:
            if action != 'KEEP':
                raise ValueError('A new row must use KEEP.')
            additions.append((text(values[leafcol+5]) or text(values[0]), text(values[leafcol]), positive(values[leafcol+1])))
            continue
        if rid not in ref or rid in seen:
            raise ValueError('BOM Map Row ID is invalid or duplicated.')
        seen.add(rid)
        before, entry = ref[rid]; chain = entry['path']; intermediate = chain[:-1] if len(chain)>1 else chain
        editable = {leafcol,leafcol+6}
        if len(chain)>1:
            editable.add(leafcol+1)
        for i, edge in enumerate(intermediate):
            editable.update((1+2*i,2+2*i))
            propose(edge,'sku',values[1+2*i] if values[1+2*i] == before[1+2*i] else text(values[1+2*i]))
            propose(edge,'quantity',positive(values[2+2*i]))
        for col, (new, old) in enumerate(zip(values, before)):
            equal = math.isclose(new, old, rel_tol=1e-12, abs_tol=1e-12) if isinstance(new, (int,float)) and isinstance(old, (int,float)) else text(new) == text(old)
            if col not in editable and not equal:
                raise ValueError(f'BOM Map: {headers[col]} is a reference field and must remain unchanged.')
        propose(chain[-1], 'sku', values[leafcol] if values[leafcol] == before[leafcol] else text(values[leafcol]))
        if len(chain)>1:
            propose(chain[-1], 'quantity', positive(values[leafcol+1]))
        if action == 'REMOVE':
            proposals[(chain[-1]['key'],chain[-1]['line']['id'],'remove')] = True
    if seen != set(ref):
        raise ValueError('BOM Map rows are missing. Keep rows and use REMOVE for component removal.')
    out = Workbook(); out.remove(out.active)
    for name in ('Products','BOMs','Components'):
        out.create_sheet(name).append(HEADERS[name])
    for name in ('Products','Non-BOM'):
        if name not in wb or wb[name].max_row > 20001:
            raise ValueError(f'{name} sheet is missing or too large.')
        rows = wb[name].iter_rows()
        if [c.value for c in next(rows)] != HEADERS['Products']+['Supplier Code']:
            raise ValueError(f'{name}: keep the original column headers.')
        for row in rows:
            if any(c.data_type == 'f' for c in row):
                raise ValueError(f'{name}: use values, not formulas.')
            if not any(c.value is not None for c in row):
                continue
            code = row[0].value
            old = record['scope']['products'].get(code, {})
            codes = record.get('supplier_codes', {}).get(code,old.get('supplier_codes',[]))
            if text(row[4].value) != '; '.join(codes):
                raise ValueError('Supplier Code is an Odoo reference field. Keep it unchanged.')
            out['Products'].append([c.value for c in row[:4]])
    if 'BOMs' not in wb or wb['BOMs'].max_row > 20001:
        raise ValueError('BOMs sheet is missing or too large.')
    if [c.value for c in next(wb['BOMs'].iter_rows())] != HEADERS['BOMs']:
        raise ValueError('BOMs: keep the original column headers.')
    for row in wb['BOMs'].iter_rows(min_row=2):
        if any(c.data_type == 'f' for c in row):
            raise ValueError('BOMs: use values, not formulas.')
        out['BOMs'].append([c.value for c in row[:4]])
    for key,bom in record['scope']['boms'].items():
        for line in bom['components']:
            prefix = (key,line['id'])
            removed = proposals.get(prefix+('remove',))
            if removed and any(prefix+(f,) in proposals for f in ('sku','quantity')):
                raise ValueError('A BOM line cannot be removed and edited in the same file.')
            out['Components'].append([key,line['id'],'REMOVE' if removed else 'KEEP',
                                      proposals.get(prefix+('sku',),line['sku']),proposals.get(prefix+('quantity',),line['quantity'])])
    for parent,code,qty in additions:
        matches = [row[0].value for row in out['BOMs'].iter_rows(min_row=2) if row[1].value == parent]
        if len(matches) != 1:
            raise ValueError('New component: Parent BOM SKU must identify exactly one BOM in BOMs.')
        out['Components'].append([matches[0],'','KEEP',code,qty])
    return out
