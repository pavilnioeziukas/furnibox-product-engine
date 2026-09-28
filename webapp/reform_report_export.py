"""Readable workbook export of a saved comparison; technical rows stay in an appendix."""
import io
from math import ceil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_report(report, view):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Santrauka'
    ws.append(['Reform versijų pakeitimai'])
    ws.append(['Ankstesnis failas', report['old']])
    ws.append(['Naujas failas', report['new']])
    ws.append(['Lyginamas lapas', report['sheet']])
    ws.append(['Parengta UTC', report.get('created_at','')])
    ws.append(['Nauji SKU iš viso', len(view['added'])])
    ws.append(['Iš jų su nauju BOM', view['with_bom']])
    ws.append(['Iš jų be savo BOM', view['without_bom']])
    ws.append(['Nauji BOM', len(view['boms'])])
    ws.append(['Pakeisti esami BOM', len(view['changed'])])
    ws.append(['Nepakitę BOM', view['unchanged']])
    ws.append(['Pašalinti SKU', len(view['removed'])])
    ws.append(['Pašalinti BOM', len(view['removed_boms'])])
    ws.append(['Kaip skaityti', 'SKU ir BOM skaičiai nesumuojami: naujas SKU gali turėti naują BOM. Naujas BOM gali būti sukurtas ir anksčiau buvusiam SKU.'])
    ws.append(['Apimtis', 'Lyginamos pasirinkto lapo išsaugotos reikšmės. Pozicijų buvimas Odoo ir pasirengimas importui netikrinami.'])
    ws.append([])
    ws.append(['Naujų SKU grupė', 'Savo BOM', 'Kiekis'])
    for group in view['sku_groups']:
        ws.append([group['category'], 'Taip' if group['has_bom'] else 'Ne', group['count']])
    ws.append([])
    bom_header = ws.max_row+1
    ws.append(['Naujų BOM grupė', 'Kodo pradžia', 'Kiekis'])
    for group in view['bom_groups']:
        ws.append([group['category'], group['code_group'], group['count']])
    summary_headers = [1,17,bom_header]

    def table(title, columns, rows, widths):
        tab = wb.create_sheet(title)
        tab.append(columns)
        for row in rows:
            tab.append(row)
        if tab.max_row == 1:
            tab.append(['Nėra'])
        tab.freeze_panes = 'B2'
        tab.auto_filter.ref = tab.dimensions
        for index, width in enumerate(widths,1):
            tab.column_dimensions[get_column_letter(index)].width = width
        tab.print_title_rows = '1:1'
        return tab

    table('Nauji SKU', ['SKU','Pavadinimas','Grupė','Savo BOM','Plotis, mm','Aukštis, mm','Gylis / storis, mm'],
          ([p['sku'],p['name'] or 'Pavadinimas nepateiktas',p['category'],'Taip' if p['has_bom'] else 'Ne',p['width'],p['height'],p['depth']] for p in view['added']), [36,72,26,14,16,16,20])
    table('Nauji BOM', ['BOM SKU','Pavadinimas','Grupė','Komponentų pozicijos'],
          ([p['sku'],p['name'] or 'Pavadinimas nepateiktas',p['category'],len(p['parts'])] for p in view['boms']), [36,80,28,24])
    table('BOM komponentai', ['BOM SKU','Komponento SKU','Kiekis','Grupė','Vieneto kaina faile'],
          ([p['sku'],c['sku'],c['quantity'],c['group'],c['price']] for p in view['boms'] for c in p['parts']), [36,36,16,30,24])
    table('Esamų BOM pakeitimai', ['BOM SKU','Kas pasikeitė'],
          ([p['sku'],'\n'.join(p['descriptions'])] for p in view['changed']), [36,110])
    table('Pašalintos pozicijos', ['Tipas','SKU','Pavadinimas','Grupė'],
          ([kind,p['sku'],p['name'] or 'Pavadinimas nepateiktas',p['category']] for kind,products in [('SKU',view['removed']),('BOM',view['removed_boms'])] for p in products), [14,36,72,28])
    # Retain the original tab name for existing integrations and full audit access.
    table('Pakeitimai', ['Techninis priedas: pakeitimas','BOM SKU','Komponentas','Laukas','Buvo','Tapo'],
          ([r[k] for k in ('status','sku','part','field','before','after')] for r in report['changes']), [32,36,36,30,48,48])
    ws.column_dimensions['A'].width=34
    ws.column_dimensions['B'].width=100
    ws.column_dimensions['C'].width=16
    ws.freeze_panes='B6'
    for tab in wb:
        tab.sheet_view.showGridLines=False
        tab.sheet_properties.pageSetUpPr.fitToPage=True
        tab.page_setup.orientation='landscape'
        tab.page_setup.paperSize=tab.PAPERSIZE_A4
        tab.page_setup.fitToWidth=1
        tab.page_setup.fitToHeight=0
        for row in tab:
            lines = 1
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str):
                    cell.data_type='s'
                width = max(8, tab.column_dimensions[cell.column_letter].width-3)
                lines = max(lines, sum(max(1,ceil(len(line)/width)) for line in str(cell.value or '').split('\n')))
            tab.row_dimensions[row[0].row].height=min(409,max(30,lines*16+10))
        for number in summary_headers if tab==ws else [1]:
            for cell in tab[number]:
                cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='204C3B')
        tab.print_options.horizontalCentered=True
    for row in (14,15):
        ws.row_dimensions[row].height=48
    output=io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
