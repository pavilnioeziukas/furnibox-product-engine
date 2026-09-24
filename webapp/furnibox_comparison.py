import csv
import io
import math
import secrets
from pathlib import Path
from datetime import datetime, timezone

from flask import Blueprint, abort, current_app, redirect, render_template, request, Response, session, url_for
import furnibox_comparison as model

comparison = Blueprint('furnibox_comparison', __name__)
COLUMNS = [
 ('sku','Detalės kodas'), ('group','Grupė'), ('furnibox_coefficient','Realus Furnibox koef.'),
 ('reform','Reform pardavimo kaina, €/vnt.'), ('purchase','Furnibox pirkimo kaina, €/vnt.'),
 ('calculated_purchase','Reali pirkimo kaina iš Furnix, €/vnt.'), ('coefficient','Furnix kategorijos koef.'),
 ('furnibox_coefficient','Realus Furnibox koef.'), ('material','Plokštė + briaunavimas, €/vnt.'),
 ('difference','Skirtumas: pirkimas − plokštė ir briauna, €/vnt.'),
 ('difference_pct','Skirtumo dalis Furnix pardavimo kainoje, %'),
 ('reform_difference','Skirtumas Reform − pirkimas, €/vnt.'), ('reform_markup','Reform antkainis nuo pirkimo, %'),
 ('basis','Pirkimo kainos pagrindas'), ('note','Duomenų pastaba'), ('po','Pirkimo užsakymas'),
 ('date','Patvirtinimo data'), ('status','Katalogo būsena'), ('name','Detalės pavadinimas'), ('category','Koeficiento kategorija')]


def excel_export(rows, rates, filters):
    # Server-side export uses the application's existing openpyxl dependency.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Detalių kainos'
    sheet.append(['Furnibox | Detalių kainų palyginimas'])
    sheet.append(['Perskaičiuotų kainų išrašas, EUR/vnt., be PVM. Koeficientus keiskite svetainėje ir eksportuokite iš naujo.'])
    sheet.append([f'Sugeneruota: {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC. Eilučių: {len(rows)}.'])
    sheet.append([f'Filtrai: {filters}'])
    sheet.append([label for _, label in COLUMNS])
    sheet['H5'] = 'Realus Furnibox koef. (pakartota)'
    for row in rows:
        sheet.append([row[key] if row[key] is not None else ('Netaikoma' if key == 'material' else None)
                      for key, _ in COLUMNS])
    for cells in sheet.iter_rows(min_row=6):
        for (key, _), cell in zip(COLUMNS, cells):
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Imported text must never execute as an Excel formula.
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '0.00%' if key in ('difference_pct', 'reform_markup') else '0.0000'
    sheet.freeze_panes = 'C6'
    if rows:
        table = Table(displayName='FurniboxKainos', ref=f'A5:T{sheet.max_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = 'A5:T5'
    for index, (key, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 42 if key in ('sku','name','note') else 24
    for r in range(1, 5):
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        sheet.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[r].height = 30
    sheet.row_dimensions[5].height = 62
    settings = workbook.create_sheet('Naudoti parametrai')
    settings.append(['Eksporto metu taikyti tarifai ir koeficientai'])
    settings.append(['Šaltinis', model.source()['source']])
    settings.append(['Parametras', 'Reikšmė'])
    for key, label in model.RATE_LABELS.items():
        settings.append([label, rates[key]])
        settings.cell(settings.max_row, 2).number_format = '0.00000'
    settings.append([])
    category_header = settings.max_row + 1
    settings.append(['Kategorija', 'Furnix koeficientas', 'Eilučių eksporte'])
    counts = {}
    for row in rows:
        if row['coefficient'] is not None:
            pair = (row['category'], row['coefficient'])
            counts[pair] = counts.get(pair, 0) + 1
    for (category, coefficient), count in sorted(counts.items()):
        settings.append([category, coefficient, count])
        settings.cell(settings.max_row, 2).number_format = '0.0000'
    settings.column_dimensions['A'].width = 48
    settings.column_dimensions['B'].width = 72
    settings.column_dimensions['C'].width = 24
    settings.freeze_panes = 'B4'
    for ws, header_rows in [(sheet, [5]), (settings, [3, category_header])]:
        ws.sheet_view.showGridLines = False
        ws.cell(1, 1).font = Font(size=18, bold=True, color='174D39')
        for r in header_rows:
            for cell in ws[r]:
                cell.fill = PatternFill('solid', fgColor='174D39')
                cell.font = Font(bold=True, color='FFFFFF')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[r].height = 62 if ws == sheet else 32
    output = io.BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename="Furnibox_kainu_palyginimas.xlsx"'})


@comparison.route('/detail-comparison', methods=['GET','POST'])
def index():
    directory = Path(current_app.config['DETAIL_CALCULATOR_PATHS']['copy']).parent / 'furnibox_comparison'
    token = session.setdefault('comparison_csrf', secrets.token_urlsafe(32))
    error = None
    if request.method == 'POST':
        if not secrets.compare_digest(request.form.get('csrf_token',''), token):
            abort(400)
        try:
            action = request.form.get('action')
            if action == 'category_coefficient':
                model.save_category_coefficient(directory, request.form.get('category'), request.form.get('coefficient'))
            elif action == 'rates':
                model.save_rates(directory, request.form)
            else:
                abort(400)
            return redirect(url_for('.index', q=request.args.get('q',''), group=request.args.get('group',''),
                                    status=request.args.get('status',''), page=request.args.get('page',1), saved='1'))
        except ValueError as exc:
            error = str(exc)
    rates, all_rows = model.results(directory)
    q = request.args.get('q','').strip().casefold()
    group, status = request.args.get('group',''), request.args.get('status','')
    rows = [r for r in all_rows if (not q or q in (r['sku']+' '+r['name']).casefold())
            and (not group or r['group']==group) and (not status or r['status']==status)]
    if request.args.get('export') == 'xlsx':
        return excel_export(rows, rates, f'paieška: {q or "visos"}; grupė: {group or "visos"}; būsena: {status or "visos"}')
    if request.args.get('export') == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output, delimiter=';')
        writer.writerow([label for key,label in COLUMNS])
        for row in rows:
            values = []
            for key, _ in COLUMNS:
                value = row[key]
                if isinstance(value, (int,float)):
                    value = f'{value*100 if key in ("difference_pct","reform_markup") else value:.6f}'.replace('.',',')
                elif value is None:
                    value = 'Netaikoma' if key=='material' else ''
                elif str(value).startswith(('=','+','-','@','\t','\r')):
                    value = "'"+str(value)
                values.append(value)
            writer.writerow(values)
        return Response('\ufeff'+output.getvalue(), mimetype='text/csv; charset=utf-8',
                        headers={'Content-Disposition':'attachment; filename="Furnibox_kainu_palyginimas.csv"'})
    try:
        page = max(1, int(request.args.get('page',1)))
    except ValueError:
        abort(400)
    pages = max(1, math.ceil(len(rows)/100)); page = min(page,pages)
    return render_template('furnibox_comparison.html', rows=rows[(page-1)*100:page*100], total=len(all_rows),
        matched=len(rows), groups=sorted({r['group'] for r in all_rows}), q=request.args.get('q',''), group=group,
        status=status, page=page, pages=pages, rates=rates, labels=model.RATE_LABELS, columns=COLUMNS,
        categories=model.category_summary(all_rows),
        csrf_token=token, error=error, saved=request.args.get('saved')=='1', source=model.source())
