"""Source comparison using existing upload history; never writes to Odoo."""
import io
import json
import secrets
import uuid
from pathlib import Path
from zipfile import BadZipFile
from datetime import datetime, timezone
from flask import Blueprint, abort, current_app, render_template, request, session, redirect, url_for, send_file
from openpyxl import Workbook
from openpyxl.styles import Font
from reform_version_diff import compare_versions

versions = Blueprint('reform_versions', __name__)


def uploads():
    root = Path(current_app.config['REFORM_VERSION_UPLOADS']).resolve()
    return sorted((p for p in root.glob('*.xlsx') if p.is_file() and not p.is_symlink()),
                  key=lambda p: (p.stat().st_mtime_ns, p.name), reverse=True)


def report_path(report_id):
    try:
        normalized = str(uuid.UUID(report_id))
    except ValueError:
        abort(404)
    if normalized != report_id:
        abort(404)
    return Path(current_app.config['REFORM_VERSION_REPORTS']) / (normalized + '.json')


@versions.route('/reform-versions', methods=['GET', 'POST'])
def index():
    files = uploads()
    token = session.setdefault('version_diff_csrf', secrets.token_urlsafe(32))
    error = None
    if request.method == 'POST':
        if not secrets.compare_digest(request.form.get('csrf_token', ''), token):
            abort(400)
        allowed = {p.name: p for p in files}
        old, new = (allowed.get(request.form.get(k)) for k in ('old', 'new'))
        if old is None or new is None or old == new:
            error = 'Pasirinkite du skirtingus įkeltus failus.'
        else:
            try:
                result = compare_versions(old, new)
                result['created_at'] = datetime.now(timezone.utc).isoformat()
                report_id = str(uuid.uuid4())
                path = report_path(report_id)
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
                return redirect(url_for('.detail', report_id=report_id))
            except (ValueError, BadZipFile, OSError) as exc:
                error = str(exc)
    root = Path(current_app.config['REFORM_VERSION_REPORTS'])
    history = []
    for p in sorted(root.glob('*.json'), key=lambda p: p.stat().st_mtime, reverse=True)[:20]:
        try:
            item = json.loads(p.read_text(encoding='utf-8'))
            history.append(dict(id=p.stem, old=item['old'], new=item['new']))
        except (ValueError, KeyError, OSError):
            continue
    return render_template('reform_versions.html', files=files, error=error, token=token, history=history)


@versions.get('/reform-versions/<report_id>')
def detail(report_id):
    path = report_path(report_id)
    if not path.is_file():
        abort(404)
    report = json.loads(path.read_text(encoding='utf-8'))
    if request.args.get('format') == 'xlsx':
        wb = Workbook()
        summary = wb.active
        summary.title = 'Santrauka'
        for key in ('old', 'new', 'sheet', 'created_at'):
            summary.append([{'old':'Ankstesnė versija','new':'Nauja versija','sheet':'Duomenų lapas','created_at':'Parengta UTC'}[key], report[key]])
        for key, value in report['counts'].items():
            summary.append([key, value])
        ws = wb.create_sheet('Pakeitimai')
        ws.append(['Pakeitimas', 'BOM SKU', 'Komponentas', 'Laukas', 'Buvo', 'Tapo'])
        for row in report['changes']:
            ws.append([row[k] for k in ('status', 'sku', 'part', 'field', 'before', 'after')])
        for title, key in [('Nauji SKU', 'added_skus'), ('Pašalinti SKU', 'removed_skus')]:
            tab = wb.create_sheet(title)
            tab.append(['SKU'])
            for sku in report[key]:
                tab.append([sku])
        for tab in wb:
            tab.freeze_panes = 'A2'
            tab.auto_filter.ref = tab.dimensions
            for column in 'ABCDEF':
                tab.column_dimensions[column].width = 36
            for cell in tab[1]:
                cell.font = Font(bold=True)
            for row in tab:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name='Reform_versiju_pakeitimai.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    query = request.args.get('q', '').strip().casefold()
    rows = [r for r in report['changes'] if not query or query in ' '.join(r.values()).casefold()]
    return render_template('reform_version_report.html', report=report, rows=rows[:1000], total=len(rows), report_id=report_id, query=request.args.get('q', ''))
