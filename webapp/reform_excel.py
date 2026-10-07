"""Versioned Excel exchange for the Reform application (no Odoo writes)."""
from webapp.time_display import cet_time

import copy
import io
import secrets
import zipfile

from flask import abort, flash, redirect, request, send_file, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from webapp.reform_workspace import (reform, db, draft, baseline, owner, now, pack, unpack,
                                    digest, changes, positive, validate_proposal, visible_work, reform_visible)

HEADERS = {
    'Products': ['Product code', 'Name', 'Category', 'Unit'],
    'BOMs': ['BOM key', 'Product code', 'BOM reference', 'Output quantity'],
    'Components': ['BOM key', 'Line ID', 'Action', 'Component code', 'Quantity'],
}
LIMIT = 5 * 1024 * 1024


def exchange_table(conn):
    conn.execute('CREATE TABLE IF NOT EXISTS file_exports (id TEXT PRIMARY KEY, owner TEXT NOT NULL, created TEXT NOT NULL, payload BLOB NOT NULL)')


def scope_for(work, sku, full=False):
    target = visible_work(work)['target']
    if full:
        return {k: target[k] for k in ('products', 'boms')}
    if sku and sku not in target['products']:
        raise ValueError('Select a product from the Reform catalogue.')
    products, boms, pending = {}, {}, [sku] if sku else []
    by_product = {}
    for key, bom in target['boms'].items():
        by_product.setdefault(bom['sku'], []).append((key, bom))
    while pending:
        code = pending.pop()
        if code in products:
            continue
        products[code] = target['products'][code]
        for key, bom in by_product.get(code, []):
            if bom.get('read_only'):
                raise ValueError('This product includes a BOM that requires Furnibox review before file editing.')
            boms[key] = bom
            pending.extend(line['sku'] for line in bom['components'])
    return {'products': products, 'boms': boms}


def options_for(work):
    products = visible_work(work)['target']['products']
    return {
        'categories': sorted({p.get('category', '') for p in products.values()} - {''}),
        'units': sorted({p['uom'] for p in products.values()}),
        'codes': sorted(k for k, p in products.items() if p['active'] and not p.get('read_only')),
    }


def make_book(token, record):
    """Runtime export uses the application's existing openpyxl dependency."""
    wb = Workbook()
    guide = wb.active
    guide.title = 'Instructions'
    for row in [
        ['REFORM BOM EXCHANGE', ''],
        ['Product', 'Full Reform catalogue' if record.get('full') else record['sku'] or 'New products and BOMs'],
        ['Read-only records', 'Records listed in Read-only must remain unchanged. Internal assembly components are excluded from this Reform view.'],
        ['Source captured (CET)', cet_time(record['work']['base']['captured_at'])],
        ['1. Edit', 'Products: edit names/categories; append new products with code, name, category and unit. Existing codes and units stay unchanged.'],
        ['2. Define BOMs', 'BOMs: edit reference or output quantity. For a new BOM, add a unique key starting NEW-, for example NEW-1.'],
        ['3. Edit components', 'Components: change code or quantity. Add rows with a blank Line ID. To remove an existing component, set Action to REMOVE.'],
        ['Keep existing rows', 'Do not delete existing rows, change BOM keys or Line IDs. Mark component removals explicitly. Row order does not matter.'],
        ['Nested BOMs', 'All editable sub-BOMs are included. BOM key identifies which assembly each component belongs to.'],
        ['Category and unit', 'Select values from the provided lists. New components may be added in Products and referenced immediately in Components.'],
        ['4. Return the file', 'Save as .xlsx. Open Excel exchange in the engine, upload the file and inspect Current / Proposed.'],
        ['5. Submit a version', 'Upload saves a draft only. Open Versions, collect all changes and submit the complete version to Furnibox. No automatic Odoo changes.'],
        ['Version check', 'The engine checks this export against its saved version. If the source or your draft changed, download a fresh file.'],
        ['Values only', 'Use values, not formulas. Keep sheet names and column headers unchanged.'],
        ['Reference catalogue', 'Catalogue lists available component codes and names. Changes to Catalogue and Lists are ignored.'],
    ]:
        guide.append(row)
    scope = record['scope']
    rows = {
        'Products': [[k, p['name'], p.get('category', ''), p['uom']] for k, p in scope['products'].items()],
        'BOMs': [[k, b['sku'], b['code'], b['quantity']] for k, b in scope['boms'].items()],
        'Components': [[k, line['id'], 'KEEP', line['sku'], line['quantity']]
                       for k, b in scope['boms'].items() for line in b['components']],
    }
    for title, headers in HEADERS.items():
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        for row in rows[title]:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    locked = wb.create_sheet('Read-only')
    locked.append(['Record type', 'Key', 'Reason'])
    for kind in ('products', 'boms'):
        for key, value in scope[kind].items():
            if value.get('read_only'):
                locked.append([kind, key, value['read_only']])
    catalogue = wb.create_sheet('Catalogue')
    catalogue.append(['Product code', 'Name', 'Category', 'Unit'])
    products = record['work']['target']['products']
    for code in record['options']['codes']:
        p = products[code]
        catalogue.append([code, p['name'], p.get('category', ''), p['uom']])
    catalogue.freeze_panes = 'A2'
    catalogue.auto_filter.ref = catalogue.dimensions
    lists = wb.create_sheet('Lists')
    lists.append(['Categories', 'Units'])
    for i in range(max(len(record['options']['categories']), len(record['options']['units']))):
        lists.append([record['options'][key][i] if i < len(record['options'][key]) else None
                      for key in ('categories', 'units')])
    for name, col, count, target in [('Categories', 'A', len(record['options']['categories']), 'C'),
                                      ('Units', 'B', len(record['options']['units']), 'D')]:
        if count:
            wb.defined_names.add(DefinedName(name, attr_text=f"'Lists'!${col}$2:${col}${count + 1}"))
            rule = DataValidation(type='list', formula1=name, allow_blank=False)
            rule.errorTitle = 'Choose an existing value'
            rule.error = 'Select a value from the list.'
            rule.showErrorMessage = True
            wb['Products'].add_data_validation(rule)
            rule.add(f'{target}2:{target}10001')
    action = DataValidation(type='list', formula1='"KEEP,REMOVE"', allow_blank=False)
    wb['Components'].add_data_validation(action)
    action.add('C2:C20001')
    meta = wb.create_sheet('_Exchange')
    meta.append(['reform-exchange-v1', token])
    meta.sheet_state = 'veryHidden'
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='21675B')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        sheet.row_dimensions[1].height = 27
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Product names are values, never spreadsheet formulas.
                cell.font = Font(name='Calibri', size=11, color='164D91' if sheet.title in HEADERS else '243B39')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 32
        widths = {'Products': [40, 55, 45, 18], 'BOMs': [22, 40, 45, 22],
                  'Components': [22, 22, 16, 44, 20], 'Catalogue': [40, 55, 45, 18],
                  'Lists': [55, 22], 'Instructions': [26, 105]}.get(sheet.title, [25, 50])
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64+i)].width = width
    for row in range(4, guide.max_row + 1):
        guide.row_dimensions[row].height = 45
    if record.get('bom_format') == 1:
        from webapp.reform_bom import add_vertical_bom
        add_vertical_bom(wb, record)
    elif record.get('map_format') == 2:
        from webapp.reform_map import add_map
        add_map(wb, record)
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return data


def read_book(content):
    if not content or len(content) > LIMIT:
        raise ValueError('Choose an .xlsx file smaller than 5 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024 or len(archive.infolist()) > 500:
                raise ValueError('The workbook is too large to process.')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError('Unable to read this workbook. Upload the original engine .xlsx template.') from exc
    return wb


def text(value):
    return '' if value is None else str(value).strip()


def zero_quantity(value):
    """Return true only for an explicit numeric zero (including Excel text such as 0,0)."""
    try:
        return float(str(value).strip().replace(',', '.')) == 0
    except (ValueError, TypeError):
        return False


def rows_from(wb, name):
    if name not in wb or wb[name].max_row > 20001 or wb[name].max_column > 20:
        raise ValueError(f'{name}: sheet is missing or exceeds the row/column limit.')
    rows = wb[name].iter_rows()
    header = next(rows)
    if [text(c.value) for c in header[:len(HEADERS[name])]] != HEADERS[name]:
        raise ValueError(f'{name}: column headers have changed. Use the original template.')
    result = []
    for number, cells in enumerate(rows, 2):
        if any(c.data_type == 'f' for c in cells):
            raise ValueError(f'{name}, row {number}: replace formulas with values.')
        values = [c.value for c in cells[:len(HEADERS[name])]]
        values += [None] * (len(HEADERS[name]) - len(values))
        if any(v is not None and text(v) for v in values):
            result.append((number, values))
    return result


def apply_book(wb, record, reserved):
    if record.get('bom_format') == 1:
        from webapp.reform_bom import translate
        translated = translate(wb, record)
        try:
            return apply_book(translated, {**record, 'bom_format': 0, 'map_format': 1}, reserved)
        finally:
            translated.close()
    if record.get('map_format') == 2:
        from webapp.reform_map import translate
        translated = translate(wb, record)
        try:
            return apply_book(translated, {**record, 'map_format': 1}, reserved)
        finally:
            translated.close()
    target = copy.deepcopy(record['work']['target'])
    scope = record['scope']
    options = record['options']
    seen = set()
    for n, row in rows_from(wb, 'Products'):
        code, name, category, unit = map(text, row)
        if row[0] in scope['products']:
            code = row[0]
        prefix = f'Products, row {n}: '
        old = scope['products'].get(code)
        if old and row[:2] == [code, old['name']] and text(row[2]) == text(old.get('category', '')) and row[3] == old['uom']:
            if code.casefold() in seen:
                raise ValueError(prefix + 'duplicate product code.')
            seen.add(code.casefold())
            continue
        if not code or len(code) > 100 or not name or len(name) > 300 or not reform_visible(code):
            raise ValueError(prefix + 'enter a valid Reform product code and name.')
        if code.casefold() in seen:
            raise ValueError(prefix + 'duplicate product code.')
        seen.add(code.casefold())
        old = scope['products'].get(code)
        if old and old.get('read_only'):
            if [name, category, unit] != [old['name'], old.get('category', ''), old['uom']]:
                raise ValueError(prefix + 'this product is read-only.')
            continue
        if not old and code.casefold() in reserved:
            raise ValueError(prefix + 'product code already exists. Use it in Components without adding a new card.')
        if not old or category != old.get('category', ''):
            if category not in options['categories']:
                raise ValueError(prefix + 'choose an existing category from Lists.')
        if unit not in options['units'] or (old and unit != old['uom']):
            raise ValueError(prefix + 'choose an existing unit; existing product units cannot be changed.')
        p = copy.deepcopy(old) if old else {'sku': code, 'active': True}
        if old and name == old['name'].strip():
            name = old['name']
        p.update(name=name, uom=unit)
        if not old or 'category' in old or category:
            p['category'] = category
        if not old:
            p['uom_id'] = next(p.get('uom_id') for p in target['products'].values() if p['uom'] == unit)
        if not old or category != old.get('category', ''):
            p['category_id'] = next((p.get('category_id') for p in target['products'].values() if p.get('category') == category), None)
        target['products'][code] = p
    if any(k.casefold() not in seen for k in scope['products']):
        raise ValueError('Products: an exported row is missing or its code changed. Keep existing rows.')
    keys, editable = {}, {}
    for n, row in rows_from(wb, 'BOMs'):
        key, code, reference = map(text, row[:3])
        if row[1] in target['products']:
            code = row[1]
        prefix = f'BOMs, row {n}: '
        if not key or key in keys or len(key) > 100:
            raise ValueError(prefix + 'BOM key is missing or duplicated.')
        old = scope['boms'].get(key)
        if old and code == old['sku'] and reference == old['code'].strip() and row[3] == old['quantity']:
            bom = copy.deepcopy(old)
            bom['components'] = []
            keys[key], editable[key] = key, bom
            continue
        if old:
            if code != old['sku']:
                raise ValueError(prefix + 'the product of an existing BOM cannot change.')
            bid = key
        else:
            if not key.startswith('NEW-'):
                raise ValueError(prefix + 'new BOM keys must start with NEW-.')
            bid = 'new-' + secrets.token_hex(8)
        p = target['products'].get(code)
        if not p or not p['active'] or p.get('read_only') or not reform_visible(code):
            raise ValueError(prefix + 'choose an active product or add its card in Products.')
        if len(reference) > 100:
            raise ValueError(prefix + 'BOM reference is too long.')
        try:
            quantity = positive(row[3])
        except ValueError as exc:
            raise ValueError(prefix + str(exc)) from exc
        bom = copy.deepcopy(old) if old else {'id': bid, 'sku': code, 'active': True, 'type': 'normal', 'uom': p['uom'], 'uom_id': p.get('uom_id')}
        if old and reference == old['code'].strip():
            reference = old['code']
        bom.update(code=reference, quantity=quantity, components=[])
        target['boms'][bid] = bom
        keys[key], editable[key] = bid, bom
    if set(scope['boms']) - set(keys):
        raise ValueError('BOMs: an exported BOM is missing or its key changed. Keep existing rows.')
    seen_lines = set()
    for n, row in rows_from(wb, 'Components'):
        key, lid, action, code = map(text, row[:4])
        if row[3] in target['products']:
            code = row[3]
        prefix = f'Components, row {n}: '
        if key not in keys or action not in ('KEEP', 'REMOVE'):
            raise ValueError(prefix + 'use a BOM key from BOMs and Action KEEP or REMOVE.')
        old = next((l for l in scope['boms'].get(key, {}).get('components', []) if l['id'] == lid), None)
        if lid:
            if not old or (key, lid) in seen_lines:
                raise ValueError(prefix + 'Line ID is invalid or duplicated for this BOM.')
            seen_lines.add((key, lid))
        # For an existing component, zero quantity is an intuitive shorthand for removal.
        # New component rows must still have a strictly positive quantity.
        if old and zero_quantity(row[4]):
            action = 'REMOVE'
        if old and action == 'KEEP' and code == old['sku'] and row[4] == old['quantity']:
            editable[key]['components'].append(copy.deepcopy(old))
            continue
        if action == 'REMOVE':
            if not old:
                raise ValueError(prefix + 'only an existing component can be marked REMOVE.')
            continue
        p = target['products'].get(code)
        if not p or not p['active'] or p.get('read_only') or not reform_visible(code):
            raise ValueError(prefix + 'component does not exist or is unavailable. Add its card in Products first.')
        try:
            quantity = positive(row[4])
        except ValueError as exc:
            raise ValueError(prefix + str(exc)) from exc
        line = copy.deepcopy(old) if old else {'id': 'new-' + secrets.token_hex(8)}
        if not old or code != old['sku']:
            line.update(uom=p['uom'], uom_id=p.get('uom_id'))
        line.update(sku=code, quantity=quantity)
        editable[key]['components'].append(line)
    expected = {(k, l['id']) for k, b in scope['boms'].items() for l in b['components']}
    if expected - seen_lines:
        raise ValueError('Components: exported rows are missing. Restore them and use REMOVE to delete a component explicitly.')
    # Row reordering is not a BOM change. Preserve the original order for existing lines.
    for key, bom in editable.items():
        positions = {l['id']: i for i, l in enumerate(scope['boms'].get(key, {}).get('components', []))}
        bom['components'].sort(key=lambda l: positions.get(l['id'], len(positions)))
        old = scope['boms'].get(key)
        if old and bom == old:
            continue  # Preserve the full stored BOM, including hidden internal rows.
        if old and old.get('read_only'):
            raise ValueError(f'BOM {key} is read-only and must remain unchanged.')
        target['boms'][keys[key]] = bom
    delta = changes(record['work']['base'], target)
    validate_proposal(target, delta)
    return target


@reform.get('/excel/download')
def excel_download():
    try:
        with db() as conn:
            work, revision = draft(conn)
            if digest(work['base']) != digest(baseline(conn)):
                raise ValueError('Your draft uses an older source. Resolve it before downloading an Excel file.')
            sku = request.args.get('product', '').strip()
            full = request.args.get('scope') == 'all'
            record = {'work': work, 'revision': revision, 'sku': sku, 'full': full, 'scope': scope_for(work, sku, full), 'options': options_for(work)}
            record['map_format'] = 1
            record['bom_format'] = 0 if request.args.get('layout') == 'legacy' else 1
            supplier_reference = conn.execute('SELECT payload FROM state WHERE id=3').fetchone()
            record['supplier_codes'] = unpack(supplier_reference['payload']) if supplier_reference else {}
            token = secrets.token_urlsafe(32)
            data = make_book(token, record)
            exchange_table(conn)
            conn.execute('INSERT INTO file_exports VALUES (?,?,?,?)', (token, owner(), now(), pack(record)))
        return send_file(data, as_attachment=True, download_name='Reform-Full-Catalogue.xlsx' if full else 'Reform-BOM.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except ValueError as exc:
        flash(str(exc))
        return redirect(url_for('reform.index', view='files'))


@reform.post('/excel/upload')
def excel_upload():
    wb = None
    try:
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.xlsx'):
            raise ValueError('Select the completed .xlsx file.')
        wb = read_book(upload.stream.read(LIMIT + 1))
        if '_Exchange' not in wb or wb['_Exchange']['A1'].value != 'reform-exchange-v1':
            raise ValueError('This is not an engine export. Download a template first.')
        token = text(wb['_Exchange']['B1'].value)
        with db() as conn:
            exchange_table(conn)
            conn.execute('BEGIN IMMEDIATE')
            export = conn.execute('SELECT * FROM file_exports WHERE id=? AND owner=?', (token, owner())).fetchone()
            if not export:
                raise ValueError('Export not found for your account. Upload a file downloaded by this account.')
            record = unpack(export['payload'])
            work, revision = draft(conn)
            if digest(record['work']['base']) != digest(baseline(conn)):
                raise ValueError('Source data has changed since this download. Download a fresh file and reapply your changes.')
            if revision != record['revision'] or digest(work) != digest(record['work']):
                raise ValueError('Your draft has changed since this download. Download a fresh file to avoid overwriting changes.')
            registry = conn.execute('SELECT payload FROM state WHERE id=2').fetchone()
            reserved = {k.casefold() for k in work['target']['products']}
            if registry:
                reserved.update(k.casefold() for k in unpack(registry['payload']))
            target = apply_book(wb, record, reserved)
            if target == work['target']:
                flash('No changes detected. Your draft is unchanged.')
            else:
                work['target'] = target
                conn.execute('INSERT OR REPLACE INTO drafts VALUES (?,?,?)', (owner(), revision+1, pack(work)))
                flash('Excel changes saved to your draft. Review Current / Proposed below, then open Versions to submit the complete version.')
        return redirect(url_for('reform.index', view='review'))
    except ValueError as exc:
        flash(str(exc))
    except Exception:
        from flask import current_app
        current_app.logger.exception('Reform Excel upload failed')
        flash('Unable to process this workbook. Keep the original sheets and upload a valid .xlsx file.')
    finally:
        if wb:
            wb.close()
    return redirect(url_for('reform.index', view='files'))

