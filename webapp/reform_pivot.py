"""Pasted, complete immediate-component BOMs; omitted parents are untouched."""
import copy
import math
import secrets
from collections import defaultdict

from openpyxl.styles import Font, PatternFill

NAME = 'Product MBOM Pivot'
ALIASES = (NAME, 'Product MBOMs - Pivot')
HEADERS = ['SKU', 'MATERIAL_SKU', 'Total']


def add_sheet(wb):
    ws = wb.create_sheet(NAME)
    ws.append(['Sum of Part Qty.'])
    ws.append(HEADERS)
    ws.freeze_panes = 'C3'
    for col, width in [('A', 38), ('B', 42), ('C', 18)]:
        ws.column_dimensions[col].width = width
    for cell in ws[2]:
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24496B')
    ws.auto_filter.ref = 'A2:C2'
    wb['Instructions'].append(['Product MBOM Pivot', 'Paste VALUES into A3:C below SKU / MATERIAL_SKU / Total (or paste the pivot with its headers at A1). Include ALL direct components of each supplied parent; absent components of that parent are removed. Omitted parents stay unchanged. Total is quantity per ONE parent unit. Do not paste a filtered/incomplete BOM, subtotal or grand total. Leave other editable sheets unchanged when using this sheet. Blank sheet is ignored.'])


def read_rows(wb):
    populated = []
    for name in ALIASES:
        if name not in wb:
            continue
        ws = wb[name]
        if ws.max_row > 20003 or ws.max_column > 3:
            raise ValueError(f'{name}: maximum 20,000 data rows and three columns.')
        rows = list(ws.iter_rows())
        header = next((i for i, row in enumerate(rows[:15])
                       if [str(c.value or '').strip() for c in row] == HEADERS), None)
        if header is None:
            raise ValueError(f'{name}: keep SKU, MATERIAL_SKU, Total headers.')
        data = []
        for n, row in enumerate(rows[header + 1:], header + 2):
            if not any(c.value is not None and str(c.value).strip() for c in row):
                continue
            if any(c.data_type == 'f' for c in row):
                raise ValueError(f'{name}, row {n}: replace formulas with values.')
            data.append((n, [c.value for c in row]))
        if data:
            populated.append(data)
    if len(populated) > 1:
        raise ValueError('Use only one Product MBOM Pivot sheet.')
    return populated[0] if populated else []


def apply_rows(rows, record):
    from webapp.reform_workspace import changes, validate_proposal, reform_visible
    target = copy.deepcopy(record['work']['target'])
    products = target['products']
    codes = {str(k).strip(): k for k in products}
    parents = defaultdict(list)
    for key, bom in record['scope']['boms'].items():
        parents[str(bom['sku']).strip()].append(key)
    grouped = defaultdict(dict)
    for n, (parent, component, value) in rows:
        parent, component = str(parent or '').strip(), str(component or '').strip()
        prefix = f'{NAME}, row {n}: '
        if not parent or not component:
            raise ValueError(prefix + 'SKU and MATERIAL_SKU are required on every row.')
        if len(parents[parent]) != 1:
            raise ValueError(prefix + 'SKU must identify exactly one existing BOM in this export.')
        code = codes.get(component)
        if code is None or not products[code]['active'] or not reform_visible(code):
            raise ValueError(prefix + 'unknown or unavailable MATERIAL_SKU. Add its product card separately first.')
        if component in grouped[parent]:
            raise ValueError(prefix + 'duplicate SKU / MATERIAL_SKU pair.')
        try:
            qty = float(str(value).strip().replace(',', '.'))
            if isinstance(value, bool) or not math.isfinite(qty) or qty < 0:
                raise ValueError()
        except (ValueError, TypeError):
            raise ValueError(prefix + 'Total must be a finite number greater than or equal to zero.')
        grouped[parent][component] = (code, qty)
    for parent, wanted in grouped.items():
        key = parents[parent][0]
        bom = target['boms'][key]
        old_totals = defaultdict(float)
        for line in bom['components']:
            old_totals[str(line['sku']).strip()] += line['quantity'] / bom['quantity']
        positive_wanted = {c: q for c, (_, q) in wanted.items() if q > 0}
        if set(positive_wanted) == set(old_totals) and all(
                math.isclose(q, old_totals[c], rel_tol=1e-10, abs_tol=1e-10)
                for c, q in positive_wanted.items()):
            continue
        if bom.get('read_only') or products[bom['sku']].get('read_only'):
            raise ValueError(f'{parent}: BOM is read-only; ask Furnibox to review it.')
        if not products[bom['sku']]['active'] or not reform_visible(bom['sku']):
            raise ValueError(f'{parent}: product is unavailable.')
        for component, (code, qty) in wanted.items():
            if products[code].get('read_only') and not math.isclose(
                    qty, old_totals.get(component, 0), rel_tol=1e-10, abs_tol=1e-10):
                raise ValueError(f'{parent} / {component}: component is read-only.')
        old_by_code = defaultdict(list)
        for line in bom['components']:
            old_by_code[str(line['sku']).strip()].append(line)
        replacement = []
        # Preserve IDs, units and ordering when a pivot aggregate is unchanged.
        for component, old in old_by_code.items():
            if component not in wanted or wanted[component][1] == 0:
                continue
            code, qty = wanted.pop(component)
            qty *= bom['quantity']
            if math.isclose(sum(x['quantity'] for x in old), qty, rel_tol=1e-10, abs_tol=1e-10):
                replacement.extend(copy.deepcopy(old))
            elif len(old) != 1:
                raise ValueError(f'{parent} / {component}: multiple existing component lines; use the BOM sheet to change them.')
            else:
                line = copy.deepcopy(old[0])
                line['quantity'] = qty
                replacement.append(line)
        for component, (code, qty) in wanted.items():
            if qty == 0 or component in old_by_code:
                continue
            p = products[code]
            replacement.append({'id': 'new-' + secrets.token_hex(8), 'sku': code,
                                'quantity': qty * bom['quantity'], 'uom': p['uom'], 'uom_id': p.get('uom_id')})
        positions = {line['id']: i for i, line in enumerate(bom['components'])}
        replacement.sort(key=lambda line: positions.get(line['id'], len(positions)))
        bom['components'] = replacement
    delta = changes(record['work']['base'], target)
    validate_proposal(target, delta)
    return target

