from openpyxl import Workbook, load_workbook
from current_reference_prices import COLUMN, current_price, add_to_workbook
from pricing_table_view import summary


def test_exact_reference_and_confirmed_conflict():
    assert current_price('SLF-PINS-HRD-4') == .4213
    assert current_price('UNI-P-ACC01-HRD019') == 5.07
    assert current_price('SKU-NOT-IN-SOURCE') is None
    assert current_price('UNI-P-ACC01-HRD019-UNKNOWN-SUFFIX') is None


def test_export_keeps_calculated_values_and_saved_snapshot(tmp_path):
    book = Workbook()
    book.active.title = 'SO LINE PRICES'
    for name in ('SO LINE PRICES', 'PRICE RESULTS'):
        sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
        sheet.append(['SKU', 'Final Reform SO Unit Price', 'Status'])
        sheet.append(['UNI-P-ACC01-HRD019', 4.9913, 'COMPLETE'])
        sheet.append(['UNKNOWN-SKU', None, 'BLOCKED'])
    add_to_workbook(book)
    add_to_workbook(book)
    path = tmp_path / 'prices.xlsx'
    book.save(path)
    saved = load_workbook(path, data_only=True)
    for name in saved.sheetnames:
        assert list(saved[name].values) == [
            ('SKU', 'Final Reform SO Unit Price', COLUMN, 'Status'),
            ('UNI-P-ACC01-HRD019', 4.9913, 5.07, 'COMPLETE'),
            ('UNKNOWN-SKU', None, None, 'BLOCKED')]
    rows = summary({'SO LINE PRICES': [{'SKU': 'UNI-P-ACC01-HRD019', COLUMN: 8.0}]})['rows']
    assert rows[0][COLUMN] == 8.0
    assert summary({'SO LINE PRICES': [{'SKU': 'UNI-P-ACC01-HRD019'}]})['rows'][0][COLUMN] == 5.07
