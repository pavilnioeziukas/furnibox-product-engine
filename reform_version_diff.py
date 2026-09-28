"""Compare source workbooks without writing to Odoo or changing active inputs."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
import re
from openpyxl import load_workbook


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_version(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next((s for s in ('BOM - Full DB', 'BOM - Input', 'BOM VERTICAL') if s in wb.sheetnames), None)
        if sheet is None:
            raise ValueError(f'{path.name}: nerastas Reform BOM duomenų lapas.')
        ws = wb[sheet]
        rows = iter(ws.iter_rows(values_only=True))
        for number, row in enumerate(rows, 1):
            labels = [text(v) for v in row]
            if 'BOM SKU Code' in labels or 'SKU Code' in labels:
                break
            if number >= 50:
                raise ValueError(f'{path.name}: nerastos BOM stulpelių antraštės.')
        else:
            raise ValueError(f'{path.name}: tuščias BOM lapas.')
        headers = {v: i for i, v in enumerate(labels) if v}
        if len(headers) != len([v for v in labels if v]):
            raise ValueError(f'{path.name}: pasikartojančios stulpelių antraštės.')
        parent_col = headers.get('BOM SKU Code', headers.get('SKU Code'))
        prefixes = [h[:-4] for h in headers if re.fullmatch(r'Part \d+ Code', h)]
        vertical = not prefixes and 'Part Code' in headers
        if vertical:
            prefixes = ['Part ']
        if not prefixes:
            raise ValueError(f'{path.name}: nerasti komponentų stulpeliai.')
        quantities = {}
        for prefix in prefixes:
            qty = next((h for h in (prefix+'Qty', prefix+'Quantity', 'Qty' if vertical else '') if h in headers), None)
            if qty is None:
                raise ValueError(f'{path.name}: trūksta {prefix}Qty stulpelio.')
            quantities[prefix] = qty
        component_headers = {h for h in headers if any(h.startswith(p) for p in prefixes)} | set(quantities.values())
        ignored = {'REF', 'BOM SKU Code', 'SKU Code'}
        metadata = set(headers) - component_headers - ignored
        products = {}
        for number, row in enumerate(rows, number+1):
            get = lambda h: text(row[headers[h]]) if headers[h] < len(row) else ''
            parent = text(row[parent_col]) if parent_col < len(row) else ''
            if not parent:
                continue
            attributes = {h: get(h) for h in metadata if get(h)}
            if parent in products and not vertical:
                raise ValueError(f'{path.name}: SKU {parent} kartojasi, eilutė {number}. Palyginimas sustabdytas.')
            product = products.setdefault(parent, {'fields': attributes, 'parts': {}})
            if product['fields'] != attributes:
                raise ValueError(f'{path.name}: SKU {parent} turi prieštaringus duomenis.')
            for prefix in prefixes:
                code = get(prefix+'Code')
                if not code:
                    continue
                try:
                    qty = Decimal(get(quantities[prefix]))
                    if not qty.is_finite() or qty <= 0:
                        raise InvalidOperation
                except InvalidOperation:
                    raise ValueError(f'{path.name}: {parent} / {code} netinkamas arba neperskaičiuotas kiekis, eilutė {number}.')
                details = {h[len(prefix):]: get(h) for h in component_headers
                           if h.startswith(prefix) and h not in (prefix+'Code', quantities[prefix]) and get(h)}
                part = product['parts'].setdefault(code, {'qty': Decimal(0), 'fields': details})
                if part['fields'] != details:
                    raise ValueError(f'{path.name}: {parent} / {code} prieštaringi komponento duomenys.')
                part['qty'] += qty
        if not products:
            raise ValueError(f'{path.name}: nėra BOM produktų; tuščia versija nelyginama.')
        return sheet, products
    finally:
        wb.close()


def compare_versions(old_path, new_path):
    old_sheet, old = read_version(old_path)
    new_sheet, new = read_version(new_path)
    if old_sheet != new_sheet:
        raise ValueError('Versijose pasirinkti skirtingi duomenų lapai. Naudokite vienodos struktūros Reform failus.')
    changes, counts = [], defaultdict(int)
    def add(status, sku, part='', field='', before='', after=''):
        changes.append(dict(status=status, sku=sku, part=part, field=field, before=text(before), after=text(after)))
    for sku in sorted(old.keys() | new.keys()):
        a, b = old.get(sku), new.get(sku)
        if a is None or b is None:
            status = 'Naujas BOM' if a is None else 'Pašalintas BOM'
            counts[status] += 1
            add(status, sku)
            for part, values in sorted((b or a)['parts'].items()):
                add(status, sku, part, 'Kiekis', values['qty'] if b is None else '', values['qty'] if a is None else '')
                for field, value in sorted(values['fields'].items()):
                    add(status, sku, part, field, value if b is None else '', value if a is None else '')
            for field, value in sorted((b or a)['fields'].items()):
                add(status, sku, field=field, before=value if b is None else '', after=value if a is None else '')
            continue
        start = len(changes)
        for field in sorted(a['fields'].keys() | b['fields'].keys()):
            av, bv = a['fields'].get(field, ''), b['fields'].get(field, '')
            if av != bv:
                add('Pakeistas laukas', sku, field=field, before=av, after=bv)
        for part in sorted(a['parts'].keys() | b['parts'].keys()):
            ap, bp = a['parts'].get(part), b['parts'].get(part)
            if ap is None or bp is None:
                add('Pridėtas komponentas' if ap is None else 'Pašalintas komponentas', sku, part, 'Kiekis', ap['qty'] if ap else '', bp['qty'] if bp else '')
                for field, value in sorted((bp or ap)['fields'].items()):
                    add('Pridėtas komponentas' if ap is None else 'Pašalintas komponentas', sku, part, field, value if bp is None else '', value if ap is None else '')
            else:
                if ap['qty'] != bp['qty']:
                    add('Pakeistas kiekis', sku, part, 'Kiekis', ap['qty'], bp['qty'])
                for field in sorted(ap['fields'].keys() | bp['fields'].keys()):
                    av, bv = ap['fields'].get(field, ''), bp['fields'].get(field, '')
                    if av != bv:
                        add('Pakeistas komponento laukas', sku, part, field, av, bv)
        counts['Pakeistas BOM' if len(changes) != start else 'Nepakitęs BOM'] += 1
    old_skus = set(old) | {p for b in old.values() for p in b['parts']}
    new_skus = set(new) | {p for b in new.values() for p in b['parts']}
    return dict(old=old_path.name, new=new_path.name, sheet=old_sheet, counts=dict(counts),
                changes=changes, added_skus=sorted(new_skus-old_skus), removed_skus=sorted(old_skus-new_skus))
